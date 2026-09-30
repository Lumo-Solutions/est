"""Drawing <-> BOQ verification end to end (API + worker body + Postgres/RLS), with the mock provider:
classes and quantities, evidence provenance, AI usage accounting, decisions (accept / reject / include /
clear) that recalculate deterministically, persistence of decisions across runs, feeding the existing
reconciliation, exports, and roles."""

from __future__ import annotations

import io
import re
import uuid

import pytest
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.ai.gateway import AIGateway
from app.ai.providers.mock import MockProvider, clear_mock_handlers, register_mock_handler
from app.ai.store import MemoryStore
from app.core.config import Settings
from app.core.context import RequestContext
from app.core.enums import Role
from app.models.takeoff import Drawing, DrawingAnnotation, DrawingMeasurement, DrawingSheet

pytestmark = pytest.mark.asyncio

TENANT = uuid.UUID("8f14e45f-ceea-4e97-8d0c-3d3b3f3c1a00")
SETTINGS = Settings(APP_DATABASE_URL="postgresql://x/y", MIGRATOR_DATABASE_URL="postgresql://x/y", S3_ENDPOINT="http://s3",
                    VLLM_API_BASE="http://vllm.internal/v1", AI_PROVIDER="mock")


@pytest.fixture(autouse=True)
def _fakes(monkeypatch):
    from app.workers.tasks import ai_reading as worker

    dispatched: list = []
    monkeypatch.setattr(worker.run_verification, "apply_async", lambda args=None, countdown=0, **kw: dispatched.append(args))
    mock = MockProvider()
    ctx = RequestContext(tenant_id=TENANT, user_id=uuid.uuid4(), sub="w", roles=frozenset({"lead_estimator"}))
    monkeypatch.setattr("app.boq.pdf_ai.default_gateway", lambda settings=None: AIGateway(settings=SETTINGS, store=MemoryStore(), provider=mock, ctx=ctx))

    def interpret(req):
        return {"lines": [{"id": i, "keywords": []} for i in re.findall(r"<<<LINE id=(\S+)", req.user)]}

    def rank(req):
        out = []
        for block in req.user.split("<<<LINE id=")[1:]:
            lid = block.split("\n", 1)[0].strip()
            desc = re.search(r"Item .*?: (.*)", block).group(1).lower()
            tokens = re.findall(r"\d{3,4}|manhole", desc)
            ranked = []
            for ref, label in re.findall(r"- ref=(\S+) \| \w+ \| (.*?) \|", block):
                hit = any(t in label.lower() for t in tokens)
                ranked.append({"ref": ref, "score": 0.9 if hit else 0.2, "selected": hit, "reason": "layer/legend agrees" if hit else "no agreement"})
            ranked.sort(key=lambda r: -r["score"])
            out.append({"id": lid, "ranked": ranked})
        return {"lines": out}

    register_mock_handler("boq_interpret", interpret)
    register_mock_handler("evidence_rank", rank)
    yield {"dispatched": dispatched, "mock": mock}
    clear_mock_handlers()


async def _seed_project(authed_client, app_engine) -> dict:
    md = authed_client(frozenset({Role.MANAGING_DIRECTOR.value}))
    pid = (await md.post("/api/v1/projects", json={"code": f"VER-{uuid.uuid4().hex[:6]}", "name": "Verification demo"})).json()["id"]
    items = {}
    head = (await md.post(f"/api/v1/projects/{pid}/boq-items", json={"item_no": "1", "description": "SEWERAGE"})).json()["id"]
    for no, desc, uom, qty in [
        ("1.1", "Supply and lay 300mm uPVC SN8 sewer pipe", "m", 245.0),
        ("1.2", "Supply and lay 150mm uPVC sewer pipe", "m", 100.0),
        ("1.3", "Precast concrete manhole 1200mm complete", "nr", 12.0),
        ("1.4", "Street lighting column 8m complete", "nr", 8.0),
        ("1.5", "Provisional sum for utilities diversion", "item", 1.0),
    ]:
        r = await md.post(f"/api/v1/projects/{pid}/boq-items", json={"parent_id": head, "item_no": no, "description": desc, "uom": uom, "boq_quantity": qty})
        items[no] = r.json()["id"]
    async with AsyncSession(app_engine) as s, s.begin():
        await s.execute(text("SELECT set_config('app.tenant_id', :t, false), set_config('app.is_system', 'on', false)"), {"t": str(TENANT)})
        d = Drawing(tenant_id=TENANT, project_id=uuid.UUID(pid), original_filename="c-101.dxf", kind="dxf", bucket="b", object_key="x",
                    size_bytes=1, sha256=uuid.uuid4().hex.ljust(64, "0"), status="ready", sheet_count=1)
        s.add(d)
        await s.flush()
        sh = DrawingSheet(tenant_id=TENANT, drawing_id=d.id, project_id=d.project_id, sheet_index=0, units="m")
        s.add(sh)
        await s.flush()
        ms = {}
        for key, (cap, kind, layer, val) in {
            "p300": ("pipe_network_topology", "pipe_length_m", "SEWER-300", 246.0),
            "p150": ("pipe_network_topology", "pipe_length_m", "SEWER-150", 150.0),
            "road": ("alignment", "alignment_length_m", "ROAD-CL", 520.0),
        }.items():
            m = DrawingMeasurement(tenant_id=TENANT, drawing_id=d.id, sheet_id=sh.id, project_id=d.project_id, capability=cap, kind=kind,
                                   value=val, unit="m", confidence=0.95, extractor_metadata={"layer": layer},
                                   bbox_min_x=0, bbox_min_y=0, bbox_max_x=10, bbox_max_y=10)
            s.add(m)
            await s.flush()
            ms[key] = str(m.id)
        a = DrawingAnnotation(tenant_id=TENANT, project_id=d.project_id, drawing_id=d.id, sheet_id=sh.id, page=1, text="12 nos precast manhole",
                              kind="count", value=12, unit="nos", bbox=[0.1, 0.2, 0.2, 0.25], source="ai", confidence=0.85)
        s.add(a)
        await s.flush()
        ms["ann"] = str(a.id)
    return {"pid": pid, "items": items, "ev": ms, "head": head}


async def _member(authed_client, app_engine, pid: str, role: str):
    """A project member with `role`; identity fixed by the first authed_client call."""
    c = authed_client(frozenset({role}))
    sub = (await c.get("/api/v1/auth/me")).json()["sub"]
    async with AsyncSession(app_engine) as s, s.begin():
        await s.execute(text("SELECT set_config('app.tenant_id', :t, false), set_config('app.is_system', 'on', false)"), {"t": str(TENANT)})
        await s.execute(text("INSERT INTO project_members (project_id, user_id, tenant_id) VALUES (:p, :u, :t)"), {"p": pid, "u": sub, "t": str(TENANT)})
    return c


async def _run(client, pid: str, role_sub_client=None) -> dict:
    import app.db.session as dbs
    from app.workers.base import build_worker_context
    from app.workers.tasks.ai_reading import _verify

    resp = await client.post(f"/api/v1/projects/{pid}/verification-runs", json={"use_ai": True})
    assert resp.status_code == 202, resp.text
    run_id = resp.json()["id"]
    dbs.get_worker_engine.cache_clear()
    sub = (await client.get("/api/v1/auth/me")).json()["sub"]
    await _verify(build_worker_context(str(TENANT), sub, ["lead_estimator"]), run_id)
    return (await client.get(f"/api/v1/verification-runs/{run_id}")).json()


def _by_item(lines):
    return {l["item_no"]: l for l in lines if l["item_no"]}


async def test_run_classes_quantities_provenance_and_usage(authed_client, app_engine, _fakes):
    seed = await _seed_project(authed_client, app_engine)
    lead = await _member(authed_client, app_engine, seed["pid"], Role.LEAD_ESTIMATOR.value)
    run = await _run(lead, seed["pid"])
    assert run["status"] == "ready" and run["progress_done"] == run["progress_total"]
    assert run["counts"] == {"match": 2, "variance": 1, "missing_in_drawings": 1, "not_verifiable": 2, "missing_in_boq": 1}
    assert run["ai_calls"] >= 2 and float(run["ai_cost_usd"]) == 0.0 and run["warnings"] == [] or run["warnings"]
    lines = (await lead.get(f"/api/v1/verification-runs/{run['id']}/lines")).json()
    assert lines["total"] == 7
    by = _by_item(lines["items"])
    l11 = by["1.1"]
    assert (l11["vclass"], l11["method"], l11["mapped_by"]) == ("match", "measured", "ai")
    assert float(l11["drawing_quantity"]) == 246.0 and abs(float(l11["delta_pct"]) - 0.408) < 0.01
    assert l11["evidence"][0]["ref"] == f"m:{seed['ev']['p300']}" and l11["evidence"][0]["page"] == 1 and l11["evidence"][0]["bbox"] is not None
    assert l11["candidates"] and l11["candidates"][0]["source"] == "ai" and l11["candidates"][0]["selected"] is True
    l12 = by["1.2"]
    assert (l12["vclass"], float(l12["delta"])) == ("variance", 50.0) and float(l12["delta_pct"]) == 50.0
    l13 = by["1.3"]
    assert (l13["vclass"], l13["method"]) == ("match", "annotated") and float(l13["confidence"]) <= 0.85
    assert by["1.4"]["vclass"] == "missing_in_drawings" and by["1.4"]["drawing_quantity"] is None and by["1.4"]["method"] == "none"
    assert by["1.5"]["vclass"] == "not_verifiable" and "Provisional" in by["1.5"]["note"]
    assert by["1"]["vclass"] == "not_verifiable"
    missing = [l for l in lines["items"] if l["vclass"] == "missing_in_boq"]
    assert len(missing) == 1 and missing[0]["boq_line_item_id"] is None and float(missing[0]["drawing_quantity"]) == 520.0
    assert missing[0]["evidence"][0]["ref"] == f"m:{seed['ev']['road']}"

    # filters and summary
    only_var = (await lead.get(f"/api/v1/verification-runs/{run['id']}/lines", params={"vclass": "variance"})).json()
    assert only_var["total"] == 1
    assert (await lead.get(f"/api/v1/verification-runs/{run['id']}/lines", params={"q": "manhole"})).json()["total"] == 1
    assert run["summary"]["top_risks"] and run["summary"]["by_trade"]


async def test_reject_accept_recalculate_and_feed_reconciliation(authed_client, app_engine):
    seed = await _seed_project(authed_client, app_engine)
    lead = await _member(authed_client, app_engine, seed["pid"], Role.LEAD_ESTIMATOR.value)
    run = await _run(lead, seed["pid"])
    lines = (await lead.get(f"/api/v1/verification-runs/{run['id']}/lines")).json()["items"]
    l12 = _by_item(lines)["1.2"]
    ref = f"m:{seed['ev']['p150']}"

    # reject: the line loses its evidence and that geometry now shows as possible missed scope
    r = await lead.post(f"/api/v1/verification-runs/{run['id']}/lines/{l12['id']}/mapping", json={"action": "reject", "evidence_ref": ref})
    assert r.status_code == 200
    body = r.json()
    assert body["line"]["vclass"] == "missing_in_drawings"
    assert body["run"]["counts"]["missing_in_drawings"] == 2 and body["run"]["counts"]["missing_in_boq"] == 2

    # accept it back: user-confirmed, confidence 1 * measurement confidence, and the EXISTING reconciliation link exists
    r = await lead.post(f"/api/v1/verification-runs/{run['id']}/lines/{l12['id']}/mapping", json={"action": "accept", "evidence_ref": ref, "note": "correct layer"})
    body = r.json()
    assert body["line"]["vclass"] == "variance" and body["line"]["mapped_by"] == "user" and float(body["line"]["confidence"]) >= 0.94
    assert body["run"]["counts"]["missing_in_boq"] == 1
    links = (await lead.get(f"/api/v1/boq-items/{seed['items']['1.2']}/measurements")).json()
    assert [m["id"] for m in links] == [seed["ev"]["p150"]]
    recon = (await lead.get(f"/api/v1/boq-items/{seed['items']['1.2']}")).json()
    assert recon["discrepancy_class"] == "variance"  # verification and reconciliation agree: one system

    # clear removes only the decision; the reconciliation link (a separate, deliberate record) is untouched
    r = await lead.post(f"/api/v1/verification-runs/{run['id']}/lines/{l12['id']}/mapping", json={"action": "clear", "evidence_ref": ref})
    assert r.json()["line"]["vclass"] == "variance"  # the AI proposal applies again
    assert len((await lead.get(f"/api/v1/boq-items/{seed['items']['1.2']}/measurements")).json()) == 1


async def test_decisions_survive_a_new_run_and_are_audited(authed_client, app_engine):
    seed = await _seed_project(authed_client, app_engine)
    lead = await _member(authed_client, app_engine, seed["pid"], Role.LEAD_ESTIMATOR.value)
    run1 = await _run(lead, seed["pid"])
    l11 = _by_item((await lead.get(f"/api/v1/verification-runs/{run1['id']}/lines")).json()["items"])["1.1"]
    ref = f"m:{seed['ev']['p300']}"
    await lead.post(f"/api/v1/verification-runs/{run1['id']}/lines/{l11['id']}/mapping", json={"action": "reject", "evidence_ref": ref})
    run2 = await _run(lead, seed["pid"])
    assert _by_item((await lead.get(f"/api/v1/verification-runs/{run2['id']}/lines")).json()["items"])["1.1"]["vclass"] == "missing_in_drawings"
    import app.db.session as dbs

    async with dbs.session_scope(RequestContext(tenant_id=TENANT, user_id=uuid.uuid4(), sub="a", roles=frozenset(), is_system=True)) as s:
        rows = (await s.execute(text("SELECT payload FROM audit_events WHERE entity_type='verification_mapping' AND project_id=:p ORDER BY seq"), {"p": seed["pid"]})).scalars().all()
    assert rows and rows[-1]["action"] == "reject" and rows[-1]["after"] == "exclude" and rows[-1]["evidence_ref"] == ref


async def test_include_arbitrary_evidence_and_dimension_guard(authed_client, app_engine):
    seed = await _seed_project(authed_client, app_engine)
    lead = await _member(authed_client, app_engine, seed["pid"], Role.LEAD_ESTIMATOR.value)
    run = await _run(lead, seed["pid"])
    l14 = _by_item((await lead.get(f"/api/v1/verification-runs/{run['id']}/lines")).json()["items"])["1.4"]
    # a length cannot be counted toward a count line
    bad = await lead.post(f"/api/v1/verification-runs/{run['id']}/lines/{l14['id']}/mapping", json={"action": "include", "evidence_ref": f"m:{seed['ev']['road']}"})
    assert bad.status_code == 422
    # the manhole annotation (12 nos) included on the lighting line: a person may, and it is 'user'-mapped
    ok = await lead.post(f"/api/v1/verification-runs/{run['id']}/lines/{l14['id']}/mapping", json={"action": "include", "evidence_ref": f"a:{seed['ev']['ann']}"})
    assert ok.status_code == 200 and ok.json()["line"]["mapped_by"] == "user" and ok.json()["line"]["vclass"] == "variance"
    unknown = await lead.post(f"/api/v1/verification-runs/{run['id']}/lines/{l14['id']}/mapping", json={"action": "include", "evidence_ref": f"m:{uuid.uuid4()}"})
    assert unknown.status_code == 404
    junk = await lead.post(f"/api/v1/verification-runs/{run['id']}/lines/{l14['id']}/mapping", json={"action": "include", "evidence_ref": "zzz"})
    assert junk.status_code == 422
    found = (await lead.get(f"/api/v1/projects/{seed['pid']}/verification-evidence", params={"q": "sewer", "kind": "measurement"})).json()["items"]
    assert len(found) == 2 and all(e["kind"] == "measurement" for e in found)


async def test_exports_xlsx_and_pdf_are_audited(authed_client, app_engine):
    from openpyxl import load_workbook

    seed = await _seed_project(authed_client, app_engine)
    lead = await _member(authed_client, app_engine, seed["pid"], Role.LEAD_ESTIMATOR.value)
    run = await _run(lead, seed["pid"])
    x = await lead.get(f"/api/v1/verification-runs/{run['id']}/export.xlsx")
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Summary", "Lines", "Missing in BOQ"]
    rows = list(wb["Lines"].iter_rows(values_only=True))
    assert rows[0][3] == "Class" and len(rows) == 1 + 6 and any(r[3] == "Variance" and r[6] == 150.0 for r in rows[1:])
    assert len(list(wb["Missing in BOQ"].iter_rows(values_only=True))) == 2
    p = await lead.get(f"/api/v1/verification-runs/{run['id']}/export.pdf")
    assert p.status_code == 200 and p.content.startswith(b"%PDF") and p.headers["content-type"] == "application/pdf"
    import app.db.session as dbs

    async with dbs.session_scope(RequestContext(tenant_id=TENANT, user_id=uuid.uuid4(), sub="a", roles=frozenset(), is_system=True)) as s:
        n = (await s.execute(text("SELECT count(*) FROM audit_events WHERE entity_type='verification_run' AND action='export' AND project_id=:p"), {"p": seed["pid"]})).scalar_one()
    assert n == 2


async def test_roles(authed_client, app_engine):
    seed = await _seed_project(authed_client, app_engine)
    pid = seed["pid"]
    # starting a run spends AI budget: estimator denied, lead allowed
    est = await _member(authed_client, app_engine, pid, Role.ESTIMATOR.value)
    assert (await est.post(f"/api/v1/projects/{pid}/verification-runs", json={})).status_code == 403
    lead = await _member(authed_client, app_engine, pid, Role.LEAD_ESTIMATOR.value)
    run = await _run(lead, pid)
    line = _by_item((await lead.get(f"/api/v1/verification-runs/{run['id']}/lines")).json()["items"])["1.2"]
    ref = f"m:{seed['ev']['p150']}"
    # working the report is estimator-level: a non-member is refused, a non-app role is refused
    outsider = authed_client(frozenset({Role.ESTIMATOR.value}))
    assert (await outsider.get(f"/api/v1/verification-runs/{run['id']}")).status_code == 404
    assert (await outsider.post(f"/api/v1/verification-runs/{run['id']}/lines/{line['id']}/mapping", json={"action": "reject", "evidence_ref": ref})).status_code in (403, 404)
    stranger = authed_client(frozenset({"platform_admin"}))
    assert (await stranger.post(f"/api/v1/verification-runs/{run['id']}/lines/{line['id']}/mapping", json={"action": "reject", "evidence_ref": ref})).status_code == 403
    # a member estimator (lowest allowed role) can accept / override
    est2 = await _member(authed_client, app_engine, pid, Role.ESTIMATOR.value)
    ok = await est2.post(f"/api/v1/verification-runs/{run['id']}/lines/{line['id']}/mapping", json={"action": "reject", "evidence_ref": ref})
    assert ok.status_code == 200
    assert (await est2.post(f"/api/v1/verification-runs/{run['id']}/lines/{line['id']}/mapping", json={"action": "bogus", "evidence_ref": ref})).status_code == 422
    assert (await est2.get(f"/api/v1/verification-runs/{uuid.uuid4()}")).status_code == 404
