"""Module D3: export into the client's original workbook, end to end --
eligibility (§3), the happy path against a real retained workbook, a
deliberately-lossy fixture blocked without accept_loss and allowed (and
audited) with it, and role gating. See docs/module-d3-plan.md."""

from __future__ import annotations

import io
import uuid
import zipfile
from datetime import datetime, timezone
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.context import RequestContext
from app.core.enums import BidSettlementStatus, QuotationExtractionMethod, QuotationLineItemStatus, QuotationStatus
from app.core.errors import ConflictError, ForbiddenError
from app.db.rls import set_rls_context
from app.models.boq import BoqImportBatch, BoqLineItem
from app.models.quotation_ingestion import Quotation, QuotationLineItem
from app.schemas.boq import BoqLineItemCreate
from app.schemas.settlement import SettlementDefaultsUpdate
from app.services import boq as boq_service
from app.services import settlement as settlement_service

pytestmark = pytest.mark.asyncio

LEAD = frozenset({"lead_estimator"})
ESTIMATOR = frozenset({"estimator"})
BD = frozenset({"bd_director"})

_MEMBER_USER_ID = uuid.UUID("00000000-0000-0000-0000-0000000000e3")


def _ctx(roles, *, tenant_id, is_system=False, user_id=None, acr=None) -> RequestContext:
    user_id = user_id or uuid.uuid4()
    return RequestContext(tenant_id=tenant_id, user_id=user_id, sub=str(user_id), roles=roles, is_system=is_system, acr=acr)


async def _system_ctx(session: AsyncSession, tenant_id: uuid.UUID) -> RequestContext:
    ctx = _ctx(frozenset(), tenant_id=tenant_id, is_system=True)
    await set_rls_context(session, ctx)
    return ctx


async def _seed_bid_submission_policy(session, tenant_id) -> None:
    from app.models.approvals import ApprovalPolicy, ApprovalPolicyTier

    policy = ApprovalPolicy(tenant_id=tenant_id, entity_type="bid_submission", name="Test policy", version=1, mode="highest_tier_only", is_active=True)
    session.add(policy)
    await session.flush()
    session.add_all(
        [
            ApprovalPolicyTier(tenant_id=tenant_id, policy_id=policy.id, seq=1, min_amount=0, max_amount=2_000_000.00, required_role="bd_director"),
            ApprovalPolicyTier(tenant_id=tenant_id, policy_id=policy.id, seq=2, min_amount=2_000_000.01, max_amount=None, max_margin_pct=8, required_role="managing_director"),
        ]
    )
    await session.flush()


async def _seed_project(session: AsyncSession) -> tuple[uuid.UUID, uuid.UUID]:
    tenant_id = uuid.uuid4()
    await session.execute(
        text("INSERT INTO tenants (id, slug, name) VALUES (:id, :slug, :name)"),
        {"id": str(tenant_id), "slug": f"acme-{uuid.uuid4().hex[:8]}", "name": "Acme"},
    )
    await _system_ctx(session, tenant_id)
    project_id = (
        await session.execute(
            text("INSERT INTO projects (tenant_id, code, name) VALUES (:t, :c, 'A') RETURNING id"),
            {"t": str(tenant_id), "c": f"D3-{uuid.uuid4().hex[:8]}"},
        )
    ).scalar_one()
    await session.execute(
        text("INSERT INTO project_members (project_id, user_id, tenant_id) VALUES (:p, :u, :t)"),
        {"p": str(project_id), "u": str(_MEMBER_USER_ID), "t": str(tenant_id)},
    )
    await _seed_bid_submission_policy(session, tenant_id)
    return tenant_id, project_id


async def _seed_accepted_quotation_line(session, ctx, *, project_id, boq_item, unit_price: str) -> QuotationLineItem:
    vendor_id = (
        await session.execute(
            text("INSERT INTO vendors (tenant_id, legal_name, normalized_name, status, primary_email) VALUES (:t, 'V', 'v', 'active', 'v@example.com') RETURNING id"),
            {"t": str(ctx.tenant_id)},
        )
    ).scalar_one()
    package_id = (
        await session.execute(
            text("INSERT INTO procurement_packages (tenant_id, project_id, name, status) VALUES (:t, :p, 'Pkg', 'sent') RETURNING id"),
            {"t": str(ctx.tenant_id), "p": str(project_id)},
        )
    ).scalar_one()
    rfq_id = (
        await session.execute(
            text("INSERT INTO rfqs (tenant_id, package_id, project_id, vendor_id, status, reply_token) VALUES (:t, :pkg, :p, :v, 'sent', :tok) RETURNING id"),
            {"t": str(ctx.tenant_id), "pkg": str(package_id), "p": str(project_id), "v": str(vendor_id), "tok": f"tok-{uuid.uuid4().hex}"},
        )
    ).scalar_one()
    quotation = Quotation(
        tenant_id=ctx.tenant_id, rfq_id=rfq_id, vendor_id=vendor_id, package_id=package_id, project_id=project_id,
        extraction_method=QuotationExtractionMethod.DETERMINISTIC_XLSX.value, version_no=1, is_current=True,
        currency="AED", vat_inclusive=True, submitted_at=datetime.now(timezone.utc), status=QuotationStatus.PROPOSED.value,
    )
    session.add(quotation)
    await session.flush()
    line_item = QuotationLineItem(
        tenant_id=ctx.tenant_id, quotation_id=quotation.id, project_id=project_id, boq_line_item_id=boq_item.id,
        unit_price=Decimal(unit_price), quantity=Decimal(str(boq_item.boq_quantity)), confidence=Decimal("1.0"),
        source="deterministic", status=QuotationLineItemStatus.ACCEPTED.value, accepted_at=datetime.now(timezone.utc),
    )
    session.add(line_item)
    await session.flush()
    return line_item


def _client_workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    ws.append(["Item No", "Description", "Unit", "Qty", "Rate", "Amount"])
    ws.append(["1", "Excavation", "m3", 250, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _inject_zip_member(data: bytes, member_name: str) -> bytes:
    buf = io.BytesIO(data)
    out = io.BytesIO()
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr(member_name, b"placeholder")
    return out.getvalue()


async def _approved_settlement_with_retained_workbook(
    rls_session, *, workbook_bytes: bytes, monkeypatch
) -> tuple[uuid.UUID, uuid.UUID, object]:
    tenant_id, project_id = await _seed_project(rls_session)
    sys_ctx = await _system_ctx(rls_session, tenant_id)
    item = await boq_service.create_line_item(
        rls_session, sys_ctx, project_id, BoqLineItemCreate(item_no="1", description="Excavation", uom="m3", boq_quantity=250.0)
    )
    await _seed_accepted_quotation_line(rls_session, sys_ctx, project_id=project_id, boq_item=item, unit_price="45.50")

    batch = BoqImportBatch(
        tenant_id=tenant_id, project_id=project_id, source_filename="tender.xlsx",
        source_object_key=f"boq-imports/{project_id}/fake.xlsx", source_sha256="0" * 64, sheet_name="BOQ",
        header_row=1, item_no_column="Item No", description_column="Description", uom_column="Unit",
        quantity_column="Qty", rate_column="E", amount_column="F", imported_by=sys_ctx.user_id,
        imported_at=datetime.now(timezone.utc),
    )
    rls_session.add(batch)
    await rls_session.flush()
    item.import_batch_id = batch.id
    item.source_row_number = 2
    await rls_session.flush()

    async def _fake_get_object_bytes(object_key, *args, **kwargs):
        assert object_key == batch.source_object_key
        return workbook_bytes

    monkeypatch.setattr(settlement_service, "get_object_bytes", _fake_get_object_bytes)

    lead_ctx = _ctx(LEAD, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, lead_ctx)
    settlement = await settlement_service.build_settlement_draft(rls_session, lead_ctx, project_id)
    await settlement_service.set_defaults(rls_session, lead_ctx, settlement.id, SettlementDefaultsUpdate(default_markup_pct=10))

    submitter = _ctx(BD, tenant_id=tenant_id, acr="silver")
    await set_rls_context(rls_session, submitter)
    await settlement_service.submit_settlement(rls_session, submitter, settlement.id)

    approver = _ctx(BD, tenant_id=tenant_id, acr="silver")
    await set_rls_context(rls_session, approver)
    approved = await settlement_service.decide_settlement(rls_session, approver, settlement.id, approve=True, note="ok")
    return tenant_id, project_id, approved


async def test_happy_path_writes_rate_cell_and_reports_ok(rls_session, monkeypatch):
    tenant_id, project_id, settlement = await _approved_settlement_with_retained_workbook(
        rls_session, workbook_bytes=_client_workbook_bytes(), monkeypatch=monkeypatch
    )
    estimator_ctx = _ctx(ESTIMATOR, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, estimator_ctx)

    report = await settlement_service.preview_original_export(rls_session, estimator_ctx, settlement.id)
    assert report["ok"] is True

    file_bytes, sha256, filename = await settlement_service.export_original_settlement(
        rls_session, estimator_ctx, settlement.id, accept_loss=False
    )
    assert filename.endswith("-original.xlsx")
    assert len(sha256) == 64
    wb = load_workbook(io.BytesIO(file_bytes))
    assert wb["BOQ"]["E2"].value == 50.05  # unit_sell_rate: 45.50 cost x 250 qty x 1.10 markup / 250


async def test_lossy_workbook_blocked_without_accept_loss_then_allowed_with_it(rls_session, monkeypatch):
    lossy = _inject_zip_member(_client_workbook_bytes(), "xl/media/image1.png")
    tenant_id, project_id, settlement = await _approved_settlement_with_retained_workbook(
        rls_session, workbook_bytes=lossy, monkeypatch=monkeypatch
    )
    estimator_ctx = _ctx(ESTIMATOR, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, estimator_ctx)

    report = await settlement_service.preview_original_export(rls_session, estimator_ctx, settlement.id)
    assert report["ok"] is False
    assert "images" in report["lost_features"]

    with pytest.raises(ConflictError):
        await settlement_service.export_original_settlement(rls_session, estimator_ctx, settlement.id, accept_loss=False)

    file_bytes, sha256, filename = await settlement_service.export_original_settlement(
        rls_session, estimator_ctx, settlement.id, accept_loss=True
    )
    assert len(file_bytes) > 0

    await _system_ctx(rls_session, tenant_id)
    audit_row = (
        await rls_session.execute(
            text("SELECT payload FROM audit_events WHERE entity_type = 'bid_settlement' AND entity_id = :e AND action = 'export' ORDER BY occurred_at DESC LIMIT 1"),
            {"e": str(settlement.id)},
        )
    ).first()
    assert audit_row[0]["accepted_loss"] is True
    assert "images" in audit_row[0]["lost_features"]


async def test_ineligible_when_lines_have_no_import_batch(rls_session):
    tenant_id, project_id = await _seed_project(rls_session)
    sys_ctx = await _system_ctx(rls_session, tenant_id)
    item = await boq_service.create_line_item(
        rls_session, sys_ctx, project_id, BoqLineItemCreate(item_no="1", description="Excavation", uom="m3", boq_quantity=250.0)
    )
    await _seed_accepted_quotation_line(rls_session, sys_ctx, project_id=project_id, boq_item=item, unit_price="45.50")
    # No BoqImportBatch/import_batch_id set at all -- e.g. manually-created BOQ item.

    lead_ctx = _ctx(LEAD, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, lead_ctx)
    settlement = await settlement_service.build_settlement_draft(rls_session, lead_ctx, project_id)
    await settlement_service.set_defaults(rls_session, lead_ctx, settlement.id, SettlementDefaultsUpdate(default_markup_pct=10))

    submitter = _ctx(BD, tenant_id=tenant_id, acr="silver")
    await set_rls_context(rls_session, submitter)
    await settlement_service.submit_settlement(rls_session, submitter, settlement.id)
    approver = _ctx(BD, tenant_id=tenant_id, acr="silver")
    await set_rls_context(rls_session, approver)
    approved = await settlement_service.decide_settlement(rls_session, approver, settlement.id, approve=True, note="ok")

    estimator_ctx = _ctx(ESTIMATOR, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, estimator_ctx)
    with pytest.raises(ConflictError):
        await settlement_service.export_original_settlement(rls_session, estimator_ctx, approved.id, accept_loss=False)


async def test_export_original_denies_role_with_no_access(rls_session, monkeypatch):
    tenant_id, project_id, settlement = await _approved_settlement_with_retained_workbook(
        rls_session, workbook_bytes=_client_workbook_bytes(), monkeypatch=monkeypatch
    )
    outsider = _ctx(frozenset(), tenant_id=tenant_id)
    await set_rls_context(rls_session, outsider)
    with pytest.raises(ForbiddenError):
        await settlement_service.export_original_settlement(rls_session, outsider, settlement.id, accept_loss=False)
