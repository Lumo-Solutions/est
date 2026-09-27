"""Module D2: generated export (state-gated, audited, leak-free -- the
leak scan itself is unit-tested against the generator directly in
tests/unit/test_settlement_export_leak.py) and win/loss capture (state-
gated, reason-code validation, role gating). See docs/module-d2-plan.md."""

from __future__ import annotations

import io
import time
import uuid
import zipfile
from datetime import datetime, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.context import RequestContext
from app.core.enums import BidSettlementStatus, QuotationExtractionMethod, QuotationLineItemStatus, QuotationStatus
from app.core.errors import ConflictError, ForbiddenError, ValidationAppError
from app.db.rls import set_rls_context
from app.models.quotation_ingestion import Quotation, QuotationLineItem
from app.models.settlement import SettlementReasonCode
from app.schemas.boq import BoqLineItemCreate
from app.schemas.settlement import ExportRequest, OutcomeRequest, SettlementDefaultsUpdate
from app.services import boq as boq_service
from app.services import settlement as settlement_service

pytestmark = pytest.mark.asyncio

LEAD = frozenset({"lead_estimator"})
ESTIMATOR = frozenset({"estimator"})
BD = frozenset({"bd_director"})

_MEMBER_USER_ID = uuid.UUID("00000000-0000-0000-0000-0000000000e2")


def _ctx(roles: frozenset[str], *, tenant_id: uuid.UUID, is_system: bool = False, user_id: uuid.UUID | None = None, acr: str | None = None) -> RequestContext:
    user_id = user_id or uuid.uuid4()
    # has_recent_step_up (fix/keycloak-step-up) requires acr AND a recent
    # auth_time, not acr alone.
    auth_time = int(time.time()) if acr else None
    return RequestContext(tenant_id=tenant_id, user_id=user_id, sub=str(user_id), roles=roles, is_system=is_system, acr=acr, auth_time=auth_time)


async def _system_ctx(session: AsyncSession, tenant_id: uuid.UUID) -> RequestContext:
    ctx = _ctx(frozenset(), tenant_id=tenant_id, is_system=True)
    await set_rls_context(session, ctx)
    return ctx


async def _seed_bid_submission_policy(session: AsyncSession, tenant_id: uuid.UUID) -> None:
    from app.models.approvals import ApprovalPolicy, ApprovalPolicyTier

    policy = ApprovalPolicy(tenant_id=tenant_id, entity_type="bid_submission", name="Test policy", version=1, mode="highest_tier_only", is_active=True)
    session.add(policy)
    await session.flush()
    session.add_all(
        [
            ApprovalPolicyTier(tenant_id=tenant_id, policy_id=policy.id, seq=1, min_amount=0, max_amount=2_000_000.00, required_role="bd_director"),
            ApprovalPolicyTier(tenant_id=tenant_id, policy_id=policy.id, seq=2, min_amount=2_000_000.01, max_amount=None, max_margin_pct=8, required_role="managing_director"),
        ]
    )
    await session.flush()


async def _seed_project(session: AsyncSession) -> tuple[uuid.UUID, uuid.UUID]:
    tenant_id = uuid.uuid4()
    await session.execute(
        text("INSERT INTO tenants (id, slug, name) VALUES (:id, :slug, :name)"),
        {"id": str(tenant_id), "slug": f"acme-{uuid.uuid4().hex[:8]}", "name": "Acme"},
    )
    await _system_ctx(session, tenant_id)
    project_id = (
        await session.execute(
            text("INSERT INTO projects (tenant_id, code, name) VALUES (:t, :c, 'A') RETURNING id"),
            {"t": str(tenant_id), "c": f"D2X-{uuid.uuid4().hex[:8]}"},
        )
    ).scalar_one()
    await session.execute(
        text("INSERT INTO project_members (project_id, user_id, tenant_id) VALUES (:p, :u, :t)"),
        {"p": str(project_id), "u": str(_MEMBER_USER_ID), "t": str(tenant_id)},
    )
    await _seed_bid_submission_policy(session, tenant_id)
    return tenant_id, project_id


async def _seed_accepted_quotation_line(session, ctx, *, project_id, boq_item, unit_price: str) -> QuotationLineItem:
    vendor_id = (
        await session.execute(
            text("INSERT INTO vendors (tenant_id, legal_name, normalized_name, status, primary_email) VALUES (:t, 'V', 'v', 'active', 'v@example.com') RETURNING id"),
            {"t": str(ctx.tenant_id)},
        )
    ).scalar_one()
    package_id = (
        await session.execute(
            text("INSERT INTO procurement_packages (tenant_id, project_id, name, status) VALUES (:t, :p, 'Pkg', 'sent') RETURNING id"),
            {"t": str(ctx.tenant_id), "p": str(project_id)},
        )
    ).scalar_one()
    rfq_id = (
        await session.execute(
            text("INSERT INTO rfqs (tenant_id, package_id, project_id, vendor_id, status, reply_token) VALUES (:t, :pkg, :p, :v, 'sent', :tok) RETURNING id"),
            {"t": str(ctx.tenant_id), "pkg": str(package_id), "p": str(project_id), "v": str(vendor_id), "tok": f"tok-{uuid.uuid4().hex}"},
        )
    ).scalar_one()
    quotation = Quotation(
        tenant_id=ctx.tenant_id, rfq_id=rfq_id, vendor_id=vendor_id, package_id=package_id, project_id=project_id,
        extraction_method=QuotationExtractionMethod.DETERMINISTIC_XLSX.value, version_no=1, is_current=True,
        currency="AED", vat_inclusive=True, submitted_at=datetime.now(timezone.utc), status=QuotationStatus.PROPOSED.value,
    )
    session.add(quotation)
    await session.flush()
    line_item = QuotationLineItem(
        tenant_id=ctx.tenant_id, quotation_id=quotation.id, project_id=project_id, boq_line_item_id=boq_item.id,
        unit_price=Decimal(unit_price), quantity=Decimal(str(boq_item.boq_quantity)), confidence=Decimal("1.0"),
        source="deterministic", status=QuotationLineItemStatus.ACCEPTED.value, accepted_at=datetime.now(timezone.utc),
    )
    session.add(line_item)
    await session.flush()
    return line_item


async def _approved_settlement(rls_session) -> tuple[uuid.UUID, uuid.UUID, object]:
    tenant_id, project_id = await _seed_project(rls_session)
    sys_ctx = await _system_ctx(rls_session, tenant_id)
    item = await boq_service.create_line_item(
        rls_session, sys_ctx, project_id, BoqLineItemCreate(item_no="1", description="Excavation", uom="m3", boq_quantity=250.0),
    )
    await _seed_accepted_quotation_line(rls_session, sys_ctx, project_id=project_id, boq_item=item, unit_price="45.50")

    lead_ctx = _ctx(LEAD, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, lead_ctx)
    settlement = await settlement_service.build_settlement_draft(rls_session, lead_ctx, project_id)
    await settlement_service.set_defaults(rls_session, lead_ctx, settlement.id, SettlementDefaultsUpdate(default_markup_pct=10))

    submitter = _ctx(BD, tenant_id=tenant_id, acr="silver")
    await set_rls_context(rls_session, submitter)
    await settlement_service.submit_settlement(rls_session, submitter, settlement.id)

    approver = _ctx(BD, tenant_id=tenant_id, acr="silver")
    await set_rls_context(rls_session, approver)
    approved = await settlement_service.decide_settlement(rls_session, approver, settlement.id, approve=True, note="ok")
    return tenant_id, project_id, approved


# --------------------------------------------------------------------------
# export
# --------------------------------------------------------------------------


async def test_export_blocked_until_approved(rls_session):
    tenant_id, project_id = await _seed_project(rls_session)
    sys_ctx = await _system_ctx(rls_session, tenant_id)
    await boq_service.create_line_item(
        rls_session, sys_ctx, project_id, BoqLineItemCreate(item_no="1", description="Excavation", uom="m3", boq_quantity=100.0)
    )
    lead_ctx = _ctx(LEAD, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, lead_ctx)
    draft = await settlement_service.build_settlement_draft(rls_session, lead_ctx, project_id)

    estimator_ctx = _ctx(ESTIMATOR, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, estimator_ctx)
    with pytest.raises(ConflictError):
        await settlement_service.export_settlement(rls_session, estimator_ctx, draft.id, ExportRequest())


async def test_export_succeeds_once_approved_and_is_audited(rls_session):
    tenant_id, project_id, settlement = await _approved_settlement(rls_session)
    estimator_ctx = _ctx(ESTIMATOR, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, estimator_ctx)

    file_bytes, sha256, filename = await settlement_service.export_settlement(rls_session, estimator_ctx, settlement.id, ExportRequest())
    assert filename.endswith(".xlsx")
    assert len(sha256) == 64

    wb = load_workbook(io.BytesIO(file_bytes))
    assert wb.sheetnames == ["BOQ"]

    # audit_events' SELECT policy restricts visibility to lead_estimator+ --
    # switch context before reading it back (estimator can export, but
    # can't read the audit trail; that's an existing, unrelated policy).
    await _system_ctx(rls_session, tenant_id)
    audit_row = (
        await rls_session.execute(
            text("SELECT action, payload FROM audit_events WHERE entity_type = 'bid_settlement' AND entity_id = :e AND action = 'export'"),
            {"e": str(settlement.id)},
        )
    ).first()
    assert audit_row is not None
    assert audit_row[1]["sha256"] == sha256


async def test_export_denies_role_with_no_settlement_access(rls_session):
    tenant_id, project_id, settlement = await _approved_settlement(rls_session)
    outsider = _ctx(frozenset(), tenant_id=tenant_id)
    await set_rls_context(rls_session, outsider)
    with pytest.raises(ForbiddenError):
        await settlement_service.export_settlement(rls_session, outsider, settlement.id, ExportRequest())


async def test_export_zip_contains_no_leak_markers_end_to_end(rls_session):
    """Real end-to-end proof (not just the unit-level generator test) that
    a settlement carrying real cost/vendor/FX data never leaks it."""
    tenant_id, project_id, settlement = await _approved_settlement(rls_session)
    estimator_ctx = _ctx(ESTIMATOR, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, estimator_ctx)

    file_bytes, _sha256, _filename = await settlement_service.export_settlement(rls_session, estimator_ctx, settlement.id, ExportRequest())
    with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
        blob = b"".join(zf.read(n) for n in zf.namelist())
    assert b"45.50" not in blob  # direct_unit_cost
    assert b"quotation_line" not in blob  # cost_source


# --------------------------------------------------------------------------
# win/loss
# --------------------------------------------------------------------------


async def test_outcome_blocked_from_draft(rls_session):
    tenant_id, project_id = await _seed_project(rls_session)
    sys_ctx = await _system_ctx(rls_session, tenant_id)
    await boq_service.create_line_item(
        rls_session, sys_ctx, project_id, BoqLineItemCreate(item_no="1", description="Excavation", uom="m3", boq_quantity=100.0)
    )
    lead_ctx = _ctx(LEAD, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, lead_ctx)
    draft = await settlement_service.build_settlement_draft(rls_session, lead_ctx, project_id)

    bd_ctx = _ctx(BD, tenant_id=tenant_id)
    await set_rls_context(rls_session, bd_ctx)
    with pytest.raises(ConflictError):
        await settlement_service.record_outcome(rls_session, bd_ctx, draft.id, OutcomeRequest(outcome="won"))


async def test_outcome_allowed_from_approved_sets_status_and_denies_lead_estimator(rls_session):
    tenant_id, project_id, settlement = await _approved_settlement(rls_session)

    lead_ctx = _ctx(LEAD, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, lead_ctx)
    with pytest.raises(ForbiddenError):
        await settlement_service.record_outcome(rls_session, lead_ctx, settlement.id, OutcomeRequest(outcome="won"))

    bd_ctx = _ctx(BD, tenant_id=tenant_id)
    await set_rls_context(rls_session, bd_ctx)
    won = await settlement_service.record_outcome(rls_session, bd_ctx, settlement.id, OutcomeRequest(outcome="won", note="great client fit"))
    assert won.status == BidSettlementStatus.WON.value
    assert won.outcome == "won"
    assert won.outcome_our_price == settlement.tender_total
    assert won.outcome_recorded_by == bd_ctx.user_id


async def test_outcome_rejects_unknown_reason_code(rls_session):
    tenant_id, project_id, settlement = await _approved_settlement(rls_session)
    bd_ctx = _ctx(BD, tenant_id=tenant_id)
    await set_rls_context(rls_session, bd_ctx)
    with pytest.raises(ValidationAppError):
        await settlement_service.record_outcome(
            rls_session, bd_ctx, settlement.id, OutcomeRequest(outcome="lost", reason_codes=["not-a-real-code"])
        )


async def test_outcome_accepts_a_seeded_reason_code(rls_session):
    tenant_id, project_id, settlement = await _approved_settlement(rls_session)
    sys_ctx = await _system_ctx(rls_session, tenant_id)
    rls_session.add(SettlementReasonCode(tenant_id=tenant_id, code="price", label="Price"))
    await rls_session.flush()

    bd_ctx = _ctx(BD, tenant_id=tenant_id)
    await set_rls_context(rls_session, bd_ctx)
    lost = await settlement_service.record_outcome(
        rls_session, bd_ctx, settlement.id, OutcomeRequest(outcome="lost", reason_codes=["price"], winning_price=Decimal("13000.00"))
    )
    assert lost.status == BidSettlementStatus.LOST.value
    assert lost.outcome_reason_codes == ["price"]
    assert lost.outcome_winning_price == Decimal("13000.00")


# --------------------------------------------------------------------------
# Phase 8f: admin CRUD for the reason-code list itself (previously
# read-only -- only app/cli.py's dev-seed ever wrote a row).
# --------------------------------------------------------------------------


async def test_create_and_update_reason_code(rls_session):
    tenant_id, _project_id, _settlement = await _approved_settlement(rls_session)
    lead_ctx = _ctx(LEAD, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, lead_ctx)

    created = await settlement_service.create_reason_code(rls_session, lead_ctx, "timeline", "Timeline")
    assert created.is_active is True

    # The win/loss form's own active-only listing excludes a deactivated
    # code; the admin screen's include_inactive=True listing still sees it.
    updated = await settlement_service.update_reason_code(rls_session, lead_ctx, created.id, is_active=False)
    assert updated.is_active is False

    active_only = await settlement_service.list_reason_codes(rls_session, lead_ctx)
    assert "timeline" not in [c.code for c in active_only]
    everything = await settlement_service.list_reason_codes(rls_session, lead_ctx, include_inactive=True)
    assert "timeline" in [c.code for c in everything]


async def test_create_reason_code_denies_estimator_role(rls_session):
    tenant_id, _project_id, _settlement = await _approved_settlement(rls_session)
    estimator_ctx = _ctx(ESTIMATOR, tenant_id=tenant_id, user_id=_MEMBER_USER_ID)
    await set_rls_context(rls_session, estimator_ctx)

    with pytest.raises(ForbiddenError):
        await settlement_service.create_reason_code(rls_session, estimator_ctx, "x", "X")
