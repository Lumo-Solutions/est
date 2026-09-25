"""Deterministic pricing-sheet parsing -- SRS changes #1 (row-id keying,
not position) and #2 (formula/non-numeric rate cells flagged, never read as
0/blank). See app/procurement/pricing_sheet_parser.py."""

from __future__ import annotations

import io
import uuid
from decimal import Decimal
from types import SimpleNamespace

from openpyxl import load_workbook

from app.procurement.content import RfqLineItem
from app.procurement.pricing_sheet import build_pricing_workbook
from app.procurement.pricing_sheet_parser import _read_rate_cell, parse_pricing_sheet

_ID_1 = uuid.uuid4()
_ID_2 = uuid.uuid4()
_ITEMS = [
    RfqLineItem(boq_line_item_id=_ID_1, item_no="1.0", description="Excavation", uom="m3", quantity=100.0),
    RfqLineItem(boq_line_item_id=_ID_2, item_no="1.1", description="Disposal", uom="m3", quantity=50.0),
]
_RFQ_ID = uuid.uuid4()
_REPLY_TOKEN = "tok-abc123"
_EXPECTED_IDS = {_ID_1, _ID_2}


def _build_reply(mutate=None) -> bytes:
    data, _sha = build_pricing_workbook(
        rfq_id=_RFQ_ID, rfq_ref="RFQ-2026-0001", reply_token=_REPLY_TOKEN, package_name="Earthworks A", items=_ITEMS,
    )
    if mutate is None:
        return data
    wb = load_workbook(io.BytesIO(data))
    mutate(wb["Pricing"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse(data: bytes):
    return parse_pricing_sheet(
        data, expected_rfq_id=_RFQ_ID, expected_reply_token=_REPLY_TOKEN, expected_boq_line_item_ids=_EXPECTED_IDS
    )


# --------------------------------------------------------------------------
# _read_rate_cell (isolated -- avoids needing to fabricate cached formula
# results inside a real saved xlsx, which openpyxl itself never computes)
# --------------------------------------------------------------------------


def _cell(value, data_type="n"):
    return SimpleNamespace(value=value, data_type=data_type)


def test_read_rate_cell_plain_number():
    result = _read_rate_cell(_cell(12.5), _cell(12.5))
    assert result.value == Decimal("12.5")
    assert not result.needs_review


def test_read_rate_cell_empty_is_not_flagged():
    result = _read_rate_cell(_cell(None), _cell(None))
    assert result.value is None
    assert not result.needs_review


def test_read_rate_cell_formula_with_cached_value_is_trusted():
    result = _read_rate_cell(_cell("=A1*2", data_type="f"), _cell(24.0))
    assert result.value == Decimal("24.0")
    assert not result.needs_review


def test_read_rate_cell_formula_without_cached_value_is_flagged_not_zero():
    result = _read_rate_cell(_cell("=A1*2", data_type="f"), _cell(None))
    assert result.value is None
    assert result.needs_review
    assert "no cached value" in result.review_reason


def test_read_rate_cell_non_numeric_text_is_flagged_not_blank():
    result = _read_rate_cell(_cell("TBC", data_type="s"), _cell("TBC"))
    assert result.value is None
    assert result.needs_review


# --------------------------------------------------------------------------
# parse_pricing_sheet
# --------------------------------------------------------------------------


def test_happy_path_reads_rates_remarks_currency_and_vat():
    def fill(ws):
        ws["E4"], ws["F4"] = 100.0, "prompt delivery"
        ws["E5"], ws["F5"] = 50.0, None
        ws["B2"], ws["E2"] = "usd", "yes"

    result = _parse(_build_reply(fill))
    assert result.ok
    assert result.currency == "USD"
    assert result.vat_inclusive is True
    by_id = {r.boq_line_item_id: r for r in result.rows}
    assert by_id[_ID_1].rate.value == Decimal("100.0")
    assert by_id[_ID_1].remarks == "prompt delivery"
    assert by_id[_ID_2].rate.value == Decimal("50.0")


def test_reordered_rows_still_parse_by_row_id():
    def swap(ws):
        row4 = [ws.cell(row=4, column=c).value for c in range(1, 8)]
        row5 = [ws.cell(row=5, column=c).value for c in range(1, 8)]
        for c, v in enumerate(row5, start=1):
            ws.cell(row=4, column=c, value=v)
        for c, v in enumerate(row4, start=1):
            ws.cell(row=5, column=c, value=v)
        ws["E4"], ws["E5"] = 50.0, 100.0  # rates follow the swapped rows

    result = _parse(_build_reply(swap))
    assert result.ok
    by_id = {r.boq_line_item_id: r for r in result.rows}
    assert by_id[_ID_1].rate.value == Decimal("100.0")
    assert by_id[_ID_2].rate.value == Decimal("50.0")


def test_deleted_row_sends_whole_sheet_to_needs_review():
    def delete_row(ws):
        ws.delete_rows(5, 1)

    result = _parse(_build_reply(delete_row))
    assert not result.ok
    assert result.rows == []


def test_duplicated_row_id_sends_whole_sheet_to_needs_review():
    def duplicate(ws):
        ws["G5"] = str(_ID_1)  # row 5 now claims the same row-id as row 4

    result = _parse(_build_reply(duplicate))
    assert not result.ok
    assert "duplicate" in result.reason


def test_inserted_row_with_unknown_row_id_sends_whole_sheet_to_needs_review():
    def insert(ws):
        ws.insert_rows(5)
        ws.cell(row=5, column=7, value=str(uuid.uuid4()))  # a row-id we never sent
        ws.cell(row=5, column=5, value=75.0)

    result = _parse(_build_reply(insert))
    assert not result.ok


def test_mismatched_meta_rfq_id_is_rejected():
    result = parse_pricing_sheet(
        _build_reply(), expected_rfq_id=uuid.uuid4(), expected_reply_token=_REPLY_TOKEN, expected_boq_line_item_ids=_EXPECTED_IDS
    )
    assert not result.ok
    assert "does not match" in result.reason


def test_mismatched_reply_token_is_rejected():
    result = parse_pricing_sheet(
        _build_reply(), expected_rfq_id=_RFQ_ID, expected_reply_token="wrong-token", expected_boq_line_item_ids=_EXPECTED_IDS
    )
    assert not result.ok


def test_not_our_workbook_at_all_is_rejected():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = "not a pricing sheet"
    buf = io.BytesIO()
    wb.save(buf)

    result = _parse(buf.getvalue())
    assert not result.ok


def test_unrecognized_vat_and_currency_text_are_none():
    def fill(ws):
        ws["B2"], ws["E2"] = "not-a-currency", "maybe"

    result = _parse(_build_reply(fill))
    assert result.ok
    assert result.currency is None
    assert result.vat_inclusive is None


def test_vat_no_is_parsed_as_false():
    def fill(ws):
        ws["E2"] = "No"

    result = _parse(_build_reply(fill))
    assert result.ok
    assert result.vat_inclusive is False
