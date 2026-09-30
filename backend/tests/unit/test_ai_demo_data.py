"""The synthetic AI demo dataset: size, tagging, determinism, and that its tender pack really produces the
deliberate matches / variances / gaps through the REAL readers and extractors (no model, no DB)."""

from __future__ import annotations

import io

from app.demo import ai_demo as demo


def test_dataset_sizes_and_tagging():
    assert 140 <= len(demo.COST_LIBRARY) <= 170
    assert len({c.code for c in demo.COST_LIBRARY}) == len(demo.COST_LIBRARY)   # unique codes
    assert all(c.code.startswith(demo.COST_CODE_PREFIX) for c in demo.COST_LIBRARY)
    assert len(demo.VENDORS) == 20 and all(v.name.startswith(demo.VENDOR_PREFIX) for v in demo.VENDORS)
    assert {"approved", "expiring", "lapsed", "pending"} <= {v.prequal for v in demo.VENDORS}
    h = demo.history()
    assert len(h) == 15 and all(t.code.startswith(demo.HISTORY_PROJECT_PREFIX) for t in h) and all(t.lines for t in h)
    assert {t.outcome for t in h} == {"won", "lost"}
    assert demo.history()[3].lines == demo.history()[3].lines  # deterministic
    assert {c.uom for c in demo.COST_LIBRARY} >= {"m", "m2", "m3", "nr"}


def test_pack_is_deterministic_and_foldered():
    a, b = demo.build_pack(), demo.build_pack()
    assert set(a) == set(b) and len(a) == 7
    assert all(p.startswith(demo.PACK_ROOT + "/") for p in a)
    assert any("/01 Drawings/" in p for p in a) and any("/02 BOQ/" in p for p in a)
    assert a[next(p for p in a if p.endswith("BOQ Rev 0.xlsx"))][:2] == b[next(p for p in b if p.endswith("BOQ Rev 0.xlsx"))][:2]


def test_excel_boq_matches_the_reference_lines():
    from openpyxl import load_workbook

    data = demo.build_pack()[f"{demo.PACK_ROOT}/02 BOQ/BOQ Rev 0.xlsx"]
    ws = load_workbook(io.BytesIO(data))["BOQ"]
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
    assert [str(r[0]) for r in rows] == [b.item_no for b in demo.BOQ]
    assert next(r for r in rows if str(r[0]) == "2.1")[3] == 245.0


def test_pdf_boq_reads_clean_pages_and_flags_the_scanned_one():
    from app.boq.pdf_reader import read_pdf_boq

    pages = read_pdf_boq(demo.build_pack()[f"{demo.PACK_ROOT}/02 BOQ/BOQ Rev 0 (print copy).pdf"])
    statuses = [p.status for p in pages]
    assert statuses[-1] == "needs_ai" and "table" in statuses
    read = {r.item_no for p in pages for r in p.rows}
    assert {"1.1", "2.1", "3.3"} <= read and "5.1" not in read   # 5.x live on the scanned page


def test_dxf_produces_the_designed_measurements_through_the_real_extractors():
    """Runs the real DXF harvester + geometry extractors on the pack drawing."""
    from types import SimpleNamespace

    from app.takeoff.dxf import index_dxf, index_dxf_geometry
    import app.takeoff.geometry.stubs  # noqa: F401  (importing it registers the extractors)
    from app.takeoff.geometry.base import get_registry

    data = demo.build_pack()[f"{demo.PACK_ROOT}/01 Drawings/C-101 Road and Drainage Plan Rev B.dxf"]
    text = " ".join(t.text_value for l in index_dxf(data) for t in l.text_entities)
    for needle in ("12 nos PRECAST MANHOLE", "ASPHALT WEARING COURSE", "DRAWING NO: C-101", "KERB L=1,040.0m"):
        assert needle in text, needle

    sheet = SimpleNamespace(units="m", scale_ratio=None)
    by_layer: dict[tuple[str, str], float] = {}
    for _layout, entities in index_dxf_geometry(data).items():
        for cap, ex in get_registry().items():
            if cap not in ("alignment", "corridor_area", "pipe_network_topology", "trench_prismoidal_volume"):
                continue  # the rest are stubs
            for m in ex.extract(sheet, entities):
                layer = (m.metadata or {}).get("layer") or ""
                by_layer[(m.kind, layer)] = by_layer.get((m.kind, layer), 0.0) + m.value
    runs = {layer: v for (kind, layer), v in by_layer.items() if kind == "pipe_run_length_m"}
    assert abs(runs["SEWER-300"] - 246.0) < 0.5 and abs(runs["SEWER-150"] - 150.0) < 0.5 and abs(runs["STORM-450"] - 80.0) < 0.5
    assert abs(next(v for (k, _l), v in by_layer.items() if k.startswith("alignment") and v > 100) - 520.0) < 0.5
    assert any(k.startswith("corridor") and abs(v - 7280.0) < 1.0 for (k, _l), v in by_layer.items())
