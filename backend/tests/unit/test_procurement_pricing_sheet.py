from __future__ import annotations

import io
import uuid

from openpyxl import load_workbook

from app.procurement.content import RfqLineItem
from app.procurement.pricing_sheet import build_pricing_workbook

_ITEMS = [
    RfqLineItem(boq_line_item_id=uuid.uuid4(), item_no="1.0", description="Excavation to formation level", uom="m3", quantity=1200.5),
    RfqLineItem(boq_line_item_id=uuid.uuid4(), item_no="1.1", description="Disposal off site", uom="m3", quantity=300.0),
]


def _build():
    rfq_id = uuid.uuid4()
    return build_pricing_workbook(
        rfq_id=rfq_id, rfq_ref="RFQ-2026-0001", reply_token="tok123", package_name="Earthworks Package A", items=_ITEMS,
    ), rfq_id


def test_pricing_sheet_contains_item_rows_and_returns_a_sha256():
    (data, sha256), _ = _build()
    assert isinstance(data, bytes) and data
    assert len(sha256) == 64

    wb = load_workbook(io.BytesIO(data))
    ws = wb["Pricing"]
    assert ws["A4"].value == "1.0"
    assert ws["B4"].value == "Excavation to formation level"
    assert ws["C4"].value == "m3"
    assert ws["D4"].value == 1200.5
    assert ws["A5"].value == "1.1"
    # Rate/Remarks columns are left blank for the vendor.
    assert ws["E4"].value is None
    assert ws["F4"].value is None


def test_sha256_is_deterministic_for_identical_content():
    rfq_id = uuid.uuid4()
    (data1, sha1) = build_pricing_workbook(
        rfq_id=rfq_id, rfq_ref="RFQ-2026-0001", reply_token="tok123", package_name="Earthworks Package A", items=_ITEMS,
    )
    (data2, sha2) = build_pricing_workbook(
        rfq_id=rfq_id, rfq_ref="RFQ-2026-0001", reply_token="tok123", package_name="Earthworks Package A", items=_ITEMS,
    )
    import hashlib

    assert sha1 == sha2 == hashlib.sha256(data1).hexdigest() == hashlib.sha256(data2).hexdigest()


def test_item_columns_are_locked_rate_and_remarks_are_not():
    (data, _sha), _ = _build()
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Pricing"]
    assert ws.protection.sheet is True
    for col in ("A", "B", "C", "D"):
        assert ws[f"{col}4"].protection.locked is True
    for col in ("E", "F"):
        assert ws[f"{col}4"].protection.locked is False


def test_row_id_column_is_hidden_locked_and_matches_boq_line_item_id():
    """SRS change #1: the deterministic parser keys off this column, not row
    position -- see app/procurement/pricing_sheet_parser.py."""
    (data, _sha), _ = _build()
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Pricing"]
    assert ws.column_dimensions["G"].hidden is True
    assert ws["G4"].value == str(_ITEMS[0].boq_line_item_id)
    assert ws["G5"].value == str(_ITEMS[1].boq_line_item_id)
    assert ws["G4"].protection.locked is True


def test_currency_and_vat_header_cells_are_present_and_editable():
    (data, _sha), _ = _build()
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Pricing"]
    assert ws["B2"].value == "AED"
    assert ws["B2"].protection.locked is False
    assert ws["E2"].value is None
    assert ws["E2"].protection.locked is False


def test_meta_sheet_is_hidden_and_maps_rows_to_boq_line_item_ids():
    (data, _sha), rfq_id = _build()
    wb = load_workbook(io.BytesIO(data))
    assert "_meta" in wb.sheetnames
    meta = wb["_meta"]
    assert meta.sheet_state == "hidden"

    values = {row[0].value: row[1].value for row in meta.iter_rows(min_row=1, max_row=3)}
    assert values["rfq_id"] == str(rfq_id)
    assert values["rfq_ref"] == "RFQ-2026-0001"
    assert values["reply_token"] == "tok123"

    mapping_rows = list(meta.iter_rows(min_row=6, values_only=True))
    assert mapping_rows[0] == (4, str(_ITEMS[0].boq_line_item_id), "1.0")
    assert mapping_rows[1] == (5, str(_ITEMS[1].boq_line_item_id), "1.1")
