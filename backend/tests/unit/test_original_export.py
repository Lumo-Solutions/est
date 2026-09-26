"""Module D3: writing settled rates into the client's own retained
workbook. Synthetic client fixtures cover the brief's required shapes --
formulas, merged cells, multiple sheets, subtotals, a hidden sheet -- plus
feature-loss detection (proven by observing an actual before/after zip
diff, not by predicting openpyxl's behaviour) and outright rejection of
.xlsm/encrypted files. See docs/module-d3-plan.md."""

from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from app.boq.original_export import RejectedOriginalError, write_settled_rates


def _client_workbook_bytes() -> bytes:
    """A synthetic client tender BOQ: item_no/description/unit/qty/rate/
    amount columns, a merged title cell, an Amount column driven by a
    formula (=D*E), a subtotal row (SUM formula), and a second, hidden
    sheet (a common real-world shape -- notes/instructions tab)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    ws.merge_cells("A1:F1")
    ws["A1"] = "Tender BOQ"
    ws.append(["Item No", "Description", "Unit", "Qty", "Rate", "Amount"])
    ws.append(["1", "Excavation", "m3", 250, None, "=D3*E3"])
    ws.append(["2", "Backfill", "m3", 100, None, "=D4*E4"])
    ws["B6"] = "Subtotal"
    ws["F6"] = "=SUM(F3:F4)"

    hidden = wb.create_sheet("Notes")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Internal notes -- not for pricing"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _inject_zip_member(data: bytes, member_name: str, member_bytes: bytes = b"placeholder") -> bytes:
    """Adds an extra zip member to an otherwise-valid workbook -- enough
    to prove the before/after diff catches its disappearance on save,
    without needing full OOXML correctness (openpyxl only round-trips
    parts it understands; anything else, referenced or not, doesn't
    survive its own save -- which is exactly the behaviour being
    detected)."""
    buf = io.BytesIO(data)
    out = io.BytesIO()
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr(member_name, member_bytes)
    return out.getvalue()


def _inject_extended_cf_marker(data: bytes) -> bytes:
    """Appends an XML comment containing the x14 conditional-formatting
    marker into the first worksheet's XML -- a comment is always valid,
    ignorable content, so the file still opens cleanly."""
    buf = io.BytesIO(data)
    out = io.BytesIO()
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            content = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"</worksheet>", b"<!-- x14:conditionalFormatting --></worksheet>")
            zout.writestr(item, content)
    return out.getvalue()


def _xlsm_bytes() -> bytes:
    """A minimal fixture standing in for a macro-enabled workbook: a
    valid xlsx zip carrying xl/vbaProject.bin, the definitive signal
    regardless of what extension the caller claims."""
    return _inject_zip_member(_client_workbook_bytes(), "xl/vbaProject.bin", b"fake vba project")


# --------------------------------------------------------------------------
# rejection
# --------------------------------------------------------------------------


def test_rejects_macro_enabled_workbook():
    with pytest.raises(RejectedOriginalError):
        write_settled_rates(
            _xlsm_bytes(), sheet_name="BOQ", rate_column="E", amount_column="F",
            writes=[(3, 45.5, 11375.0)],
        )


def test_rejects_non_zip_data_as_encrypted():
    not_a_zip = b"this is not a valid ooxml/zip file at all"
    with pytest.raises(RejectedOriginalError):
        write_settled_rates(not_a_zip, sheet_name="BOQ", rate_column="E", amount_column="F", writes=[(3, 45.5, 11375.0)])


# --------------------------------------------------------------------------
# writing + fidelity
# --------------------------------------------------------------------------


def test_writes_only_rate_cell_when_amount_is_a_formula():
    original = _client_workbook_bytes()
    output_bytes, report = write_settled_rates(
        original, sheet_name="BOQ", rate_column="E", amount_column="F",
        writes=[(3, 45.5, 11375.0), (4, 12.0, 1200.0)],
    )
    assert report.ok is True
    assert report.lost_features == []

    wb = load_workbook(io.BytesIO(output_bytes), data_only=False)
    ws = wb["BOQ"]
    assert ws["E3"].value == 45.5
    assert ws["E4"].value == 12.0
    # Amount was a formula -- never overwritten, still the original formula.
    assert ws["F3"].value == "=D3*E3"
    assert ws["F4"].value == "=D4*E4"
    assert ws["F6"].value == "=SUM(F3:F4)"


def test_writes_amount_cell_when_not_a_formula():
    original = _client_workbook_bytes()
    # Replace the formula with a plain blank cell first (a template that
    # genuinely has no formula in the Amount column).
    wb = load_workbook(io.BytesIO(original))
    wb["BOQ"]["F3"] = None
    buf = io.BytesIO()
    wb.save(buf)
    original_plain = buf.getvalue()

    output_bytes, report = write_settled_rates(
        original_plain, sheet_name="BOQ", rate_column="E", amount_column="F", writes=[(3, 45.5, 11375.0)],
    )
    assert report.ok is True
    out_ws = load_workbook(io.BytesIO(output_bytes))["BOQ"]
    assert out_ws["F3"].value == 11375.0


def test_merged_cells_sheet_names_and_hidden_sheet_survive():
    original = _client_workbook_bytes()
    output_bytes, report = write_settled_rates(
        original, sheet_name="BOQ", rate_column="E", amount_column="F", writes=[(3, 45.5, 11375.0)],
    )
    assert report.ok is True

    wb_orig = load_workbook(io.BytesIO(original))
    wb_out = load_workbook(io.BytesIO(output_bytes))
    assert wb_out.sheetnames == wb_orig.sheetnames == ["BOQ", "Notes"]
    assert {str(r) for r in wb_out["BOQ"].merged_cells.ranges} == {str(r) for r in wb_orig["BOQ"].merged_cells.ranges}
    assert wb_out["Notes"].sheet_state == "hidden"
    assert wb_out["Notes"]["A1"].value == "Internal notes -- not for pricing"


def test_unwritten_cells_outside_the_written_set_are_flagged_if_changed():
    """Proves the checker itself works by constructing a "bad" output that
    changed something it shouldn't have -- not exercising write_settled_rates
    (which never does this), but the report-building logic it shares."""
    from app.boq.original_export import _build_fidelity_report

    original = _client_workbook_bytes()
    wb = load_workbook(io.BytesIO(original))
    wb["BOQ"]["B3"] = "Excavation (tampered)"  # a description cell, never in any written set
    buf = io.BytesIO()
    wb.save(buf)
    tampered = buf.getvalue()

    report = _build_fidelity_report(original, tampered, written_coords=set(), before_features={})
    assert report.ok is False
    assert "cell_values_outside_written_set" in report.lost_features
    assert "BOQ!B3" in report.unexpected_cell_changes


# --------------------------------------------------------------------------
# feature-loss detection (observed, not predicted)
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "member_path,category",
    [
        ("xl/media/image1.png", "images"),
        ("xl/charts/chart1.xml", "charts"),
        ("xl/pivotTables/pivotTable1.xml", "pivot_tables"),
        ("xl/externalLinks/externalLink1.xml", "external_links"),
    ],
)
def test_detects_a_lost_feature_category(member_path, category):
    original = _inject_zip_member(_client_workbook_bytes(), member_path)
    # Sanity: openpyxl can still open a workbook carrying an orphan part.
    load_workbook(io.BytesIO(original))

    _output_bytes, report = write_settled_rates(
        original, sheet_name="BOQ", rate_column="E", amount_column="F", writes=[(3, 45.5, 11375.0)],
    )
    assert report.ok is False
    assert category in report.lost_features


def test_detects_extended_conditional_formatting_marker():
    original = _inject_extended_cf_marker(_client_workbook_bytes())
    load_workbook(io.BytesIO(original))  # still opens fine (a comment is harmless)

    _output_bytes, report = write_settled_rates(
        original, sheet_name="BOQ", rate_column="E", amount_column="F", writes=[(3, 45.5, 11375.0)],
    )
    assert report.ok is False
    assert "extended_conditional_formatting_or_validation" in report.lost_features


def test_clean_workbook_with_none_of_these_features_reports_ok():
    original = _client_workbook_bytes()
    _output_bytes, report = write_settled_rates(
        original, sheet_name="BOQ", rate_column="E", amount_column="F", writes=[(3, 45.5, 11375.0), (4, 12.0, 1200.0)],
    )
    assert report.ok is True
    assert report.lost_features == []
