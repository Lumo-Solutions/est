"""BOQ import: reading a named sheet (tender intake pre-fills this) instead of only the active sheet."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from app.boq.import_parser import BoqImportColumnMapping, parse_boq_file
from app.core.errors import ValidationAppError


def _workbook() -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "Project summary"
    boq = wb.create_sheet("BOQ")
    boq.append(["Client logo"])
    boq.append([])
    boq.append(["Item No", "Description", "Unit", "Qty", "Rate", "Amount"])
    boq.append(["1", "Earthworks", None, None, None, None])
    boq.append(["1.1", "Excavation", "m3", 500, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _mapping(**kw) -> BoqImportColumnMapping:
    return BoqImportColumnMapping(item_no_column="Item No", description_column="Description", uom_column="Unit", quantity_column="Qty", header_row=3, **kw)


def test_named_sheet_is_read_with_its_header_row():
    result = parse_boq_file(_workbook(), "b.xlsx", _mapping(sheet_name="BOQ"))
    assert result.error_count == 0 and [r.item_no for r in result.rows] == ["1", "1.1"]
    assert result.rows[1].boq_quantity == 500 and result.rows[1].parent_item_no == "1"


def test_without_a_sheet_name_the_active_sheet_is_used_as_before():
    result = parse_boq_file(_workbook(), "b.xlsx", _mapping())
    assert result.rows == []  # the active sheet is Summary, which has no BOQ rows below row 3
    wb = Workbook()
    wb.active.append(["Item No", "Description"])
    wb.active.append(["1", "Only sheet"])
    buf = io.BytesIO()
    wb.save(buf)
    plain = parse_boq_file(buf.getvalue(), "b.xlsx", BoqImportColumnMapping(item_no_column="Item No", description_column="Description"))
    assert [r.description for r in plain.rows] == ["Only sheet"]


def test_unknown_sheet_is_a_clear_validation_error():
    with pytest.raises(ValidationAppError, match="Sheet 'Nope' not found"):
        parse_boq_file(_workbook(), "b.xlsx", _mapping(sheet_name="Nope"))
