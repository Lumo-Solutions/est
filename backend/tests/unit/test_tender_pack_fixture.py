"""Unit tests for the synthetic tender pack fixture (no DB)."""

import hashlib
import io
import zipfile

import pytest

from tests.fixtures.tender_pack import GROUND_TRUTH, build_pack, build_zip_member_files


@pytest.fixture(scope="module")
def pack():
    return {f.name: f for f in build_pack()}


def test_deterministic():
    a = {f.name: hashlib.sha256(f.data).hexdigest() for f in build_pack()}
    b = {f.name: hashlib.sha256(f.data).hexdigest() for f in build_pack()}
    assert a == b
    ma = [hashlib.sha256(f.data).hexdigest() for f in build_zip_member_files()]
    mb = [hashlib.sha256(f.data).hexdigest() for f in build_zip_member_files()]
    assert ma == mb


def test_names_unique():
    names = [f.name for f in build_pack()]
    assert len(names) == len(set(names))
    members = [f.name for f in build_zip_member_files()]
    assert len(members) == len(set(members))
    assert not set(names) & set(members)


def test_dxf_headers(pack):
    for name in ("C-101 Rev A.dxf", "C-101 Rev B.dxf", "A-201 Rev A.dxf"):
        data = pack[name].data
        assert b"SECTION" in data[:300]
        assert b"EOF" in data[-50:]
        assert b"\r\n" in data
        assert b"\n" not in data.replace(b"\r\n", b"")
    assert b"DRAWING NO: C-101" in pack["C-101 Rev A.dxf"].data
    assert b"REV: B" in pack["C-101 Rev B.dxf"].data


def test_duplicate_bytes(pack):
    dup = pack["C-101 Rev A (copy).dxf"]
    assert dup.data == pack["C-101 Rev A.dxf"].data
    assert dup.expected_status == "duplicate"
    assert dup.expected_class == "drawing"


def test_xlsm_has_vba(pack):
    f = pack["BOQ_macro.xlsm"]
    with zipfile.ZipFile(io.BytesIO(f.data)) as zf:
        assert "xl/vbaProject.bin" in zf.namelist()
    assert (f.expected_status, f.reject_reason, f.expected_class) == ("rejected", "macro", None)


def test_zip_bomb(pack):
    f = pack["bomb.zip"]
    assert len(f.data) < 1024 * 1024
    with zipfile.ZipFile(io.BytesIO(f.data)) as zf:
        infos = zf.infolist()
        assert len(infos) == 1
        assert infos[0].file_size >= 100 * 1024 * 1024
    assert (f.expected_status, f.reject_reason) == ("rejected", "zip_bomb")


def test_pdfs_open_and_text_layers(pack):
    from pypdf import PdfReader

    pdfs = [f for f in list(pack.values()) + build_zip_member_files() if f.name.lower().endswith(".pdf")]
    assert pdfs
    for f in pdfs:
        reader = PdfReader(io.BytesIO(f.data))
        assert len(reader.pages) == 1
        text = (reader.pages[0].extract_text() or "").strip()
        if f.name == "Scan_0042.pdf":
            assert text == ""
        else:
            assert text, f.name
    add = PdfReader(io.BytesIO(pack["ADD-01.pdf"].data)).pages[0].extract_text()
    assert "ADDENDUM No. 1" in add
    s201 = PdfReader(io.BytesIO(pack["S-201 Rev B.pdf"].data)).pages[0].extract_text()
    assert "S-201" in s201 and "REV B" in s201


def test_boq_header_row(pack):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(pack["BOQ_Rev0.xlsx"].data))
    assert wb.sheetnames == ["Summary", "BOQ"]
    row3 = [c.value for c in wb["BOQ"][3]]
    assert row3 == ["Item No", "Description", "Unit", "Qty", "Rate", "Amount"]
    assert wb["BOQ"].max_row == 3 + 8


def test_transmittal_zip_members(pack):
    expected = [
        "Drawings/S-201 Rev A.pdf",
        "Forms/Form_of_Tender.pdf",
        "Correspondence/Letter_from_Client.pdf",
        "Geotech/Geotechnical_Report.pdf",
    ]
    with zipfile.ZipFile(io.BytesIO(pack["Transmittal_2.zip"].data)) as zf:
        assert sorted(zf.namelist()) == sorted(expected)
    assert pack["Transmittal_2.zip"].expected_status == "archive"
    assert sorted(f.name for f in build_zip_member_files()) == sorted(expected)


def test_ground_truth_covers_classifiable():
    expected = {
        f.name: f.expected_class
        for f in build_pack() + build_zip_member_files()
        if f.expected_status in ("classifiable", "duplicate")
    }
    assert GROUND_TRUTH == expected
    assert "bomb.zip" not in GROUND_TRUTH
    assert "Transmittal_2.zip" not in GROUND_TRUTH
    assert "BOQ_macro.xlsm" not in GROUND_TRUTH
    assert GROUND_TRUTH["Scan_0042.pdf"] == "drawing"
