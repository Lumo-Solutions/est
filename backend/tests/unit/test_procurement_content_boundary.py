"""SRS Module C1 change #7: RFQ content (email body + xlsx) may contain only
item_no/description/unit/quantity -- never a rate, budget, estimate, or any
vendor's name other than the one being addressed. See
app/procurement/content.py and docs/procurement-rfq.md."""

from __future__ import annotations

import dataclasses
import io
import uuid

from openpyxl import load_workbook

from app.procurement.content import RfqLineItem
from app.procurement.pricing_sheet import build_pricing_workbook
from app.procurement.rendering import render_rfq_email

_ADDRESSED_VENDOR = "Al Futtaim Earthworks LLC"
_OTHER_VENDOR = "Confidential Rival Contracting Co"
_FORBIDDEN_SUBSTRINGS = [_OTHER_VENDOR, "AED 1,250,000", "budget", "estimate", "internal margin"]

_ITEMS = [
    RfqLineItem(boq_line_item_id=uuid.uuid4(), item_no="2.0", description="Trench excavation 0-2m", uom="m3", quantity=450.0),
]


def test_rfq_line_item_dataclass_only_carries_the_allowed_fields():
    field_names = {f.name for f in dataclasses.fields(RfqLineItem)}
    assert field_names == {"boq_line_item_id", "item_no", "description", "uom", "quantity"}


def test_email_body_contains_no_rate_or_other_vendor_data():
    html = render_rfq_email(
        vendor_name=_ADDRESSED_VENDOR, rfq_ref="RFQ-2026-0007", package_name="Earthworks Package A",
        items=_ITEMS, due_at="2026-03-01", sender_name="INSTALLTEC Procurement Team",
    )
    assert _ADDRESSED_VENDOR in html
    for forbidden in _FORBIDDEN_SUBSTRINGS:
        assert forbidden not in html, f"{forbidden!r} leaked into the RFQ email body"


def test_xlsx_contains_no_rate_or_other_vendor_data():
    data, _sha256 = build_pricing_workbook(
        rfq_id=uuid.uuid4(), rfq_ref="RFQ-2026-0007", reply_token="tok", package_name="Earthworks Package A",
        items=_ITEMS,
    )
    wb = load_workbook(io.BytesIO(data))
    all_text = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            all_text.extend(str(v) for v in row if v is not None)
    blob = "\n".join(all_text)
    for forbidden in _FORBIDDEN_SUBSTRINGS:
        assert forbidden not in blob, f"{forbidden!r} leaked into the RFQ pricing sheet"
