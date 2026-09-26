"""Module D3: writing settled rates back into the client's own retained
tender workbook. Pure logic -- no DB/S3 -- so tests exercise it directly
against synthetic fixtures; the async orchestration (fetching the
retained original, settlement lines, audit) lives in
app/services/settlement.py::export_original_settlement. See
docs/module-d3-plan.md.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field

from openpyxl import load_workbook

_VBA_PROJECT_MEMBER = "xl/vbaProject.bin"

# Zip member prefixes that prove a feature category is present -- checked
# before writing and again after, so a loss is *observed* for this exact
# file, never merely predicted from openpyxl's documented limitations.
_FEATURE_PREFIXES: dict[str, tuple[str, ...]] = {
    "images": ("xl/media/",),
    "charts": ("xl/charts/",),
    "pivot_tables": ("xl/pivotTables/", "xl/pivotCache/"),
    "external_links": ("xl/externalLinks/",),
}


class RejectedOriginalError(Exception):
    """.xlsm / macro-enabled / encrypted -- refused outright, before any
    fidelity check even runs."""


@dataclass(frozen=True, slots=True)
class FidelityReport:
    ok: bool
    lost_features: list[str] = field(default_factory=list)
    unexpected_cell_changes: list[str] = field(default_factory=list)


def assert_not_macro_or_encrypted(data: bytes) -> None:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise RejectedOriginalError(
            "The retained original could not be opened as a workbook -- it may be password-protected/encrypted."
        )
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        if _VBA_PROJECT_MEMBER in zf.namelist():
            raise RejectedOriginalError("The retained original contains macros (.xlsm) -- rejected.")


def _feature_members(data: bytes) -> dict[str, set[str]]:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = zf.namelist()
        found = {
            category: {n for n in names if any(n.startswith(p) for p in prefixes)}
            for category, prefixes in _FEATURE_PREFIXES.items()
        }
        # Extended (x14-namespaced) conditional formatting / data
        # validation -- openpyxl's classic parser reads the base rule
        # types reliably; these extras are the ones known not to
        # round-trip. Detected by grepping each worksheet's raw XML for
        # the namespace prefix, not by trying to fully parse extLst.
        extended = set()
        for name in names:
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                content = zf.read(name).decode("utf-8", errors="ignore")
                if "x14:conditionalFormatting" in content or "x14:dataValidation" in content:
                    extended.add(name)
        found["extended_conditional_formatting_or_validation"] = extended
    return found


def write_settled_rates(
    original_bytes: bytes, *, sheet_name: str | None, rate_column: str, amount_column: str | None,
    writes: list[tuple[int, float, float]],
) -> tuple[bytes, FidelityReport]:
    """writes: (spreadsheet row, unit_sell_rate, line_amount) triples.
    Returns (output_bytes, fidelity_report) -- the caller decides what to
    do with a report that isn't ok (§6/§7 of the plan: refuse unless the
    caller explicitly accepts the loss)."""
    assert_not_macro_or_encrypted(original_bytes)
    before_features = _feature_members(original_bytes)

    wb = load_workbook(io.BytesIO(original_bytes), data_only=False)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    written_coords: set[tuple[str, str]] = set()
    for row, unit_rate, line_amount in writes:
        rate_coord = f"{rate_column}{row}"
        ws[rate_coord] = float(unit_rate)
        written_coords.add((ws.title, rate_coord))
        if amount_column:
            amount_coord = f"{amount_column}{row}"
            cell = ws[amount_coord]
            if cell.data_type != "f":  # brief's exact rule: never overwrite a formula
                cell.value = float(line_amount)
                written_coords.add((ws.title, amount_coord))

    buf = io.BytesIO()
    wb.save(buf)
    output_bytes = buf.getvalue()

    report = _build_fidelity_report(original_bytes, output_bytes, written_coords, before_features)
    return output_bytes, report


def _build_fidelity_report(
    original_bytes: bytes, output_bytes: bytes, written_coords: set[tuple[str, str]], before_features: dict[str, set[str]]
) -> FidelityReport:
    lost: list[str] = []

    after_features = _feature_members(output_bytes)
    for category, before_set in before_features.items():
        if before_set and not before_set.issubset(after_features.get(category, set())):
            lost.append(category)

    wb_orig = load_workbook(io.BytesIO(original_bytes), data_only=False)
    wb_out = load_workbook(io.BytesIO(output_bytes), data_only=False)

    if wb_orig.sheetnames != wb_out.sheetnames:
        lost.append("sheet_names_or_order")

    unexpected_cell_changes: list[str] = []
    merged_ranges_changed = False
    for name in wb_orig.sheetnames:
        if name not in wb_out.sheetnames:
            continue
        ws_orig, ws_out = wb_orig[name], wb_out[name]
        if {str(r) for r in ws_orig.merged_cells.ranges} != {str(r) for r in ws_out.merged_cells.ranges}:
            merged_ranges_changed = True

        max_row = max(ws_orig.max_row, ws_out.max_row)
        max_col = max(ws_orig.max_column, ws_out.max_column)
        for row in ws_orig.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
            for cell in row:
                coord = cell.coordinate
                if (name, coord) in written_coords:
                    continue
                if cell.value != ws_out[coord].value:
                    unexpected_cell_changes.append(f"{name}!{coord}")

    if merged_ranges_changed:
        lost.append("merged_ranges")
    if unexpected_cell_changes:
        lost.append("cell_values_outside_written_set")

    return FidelityReport(ok=not lost, lost_features=lost, unexpected_cell_changes=unexpected_cell_changes)
