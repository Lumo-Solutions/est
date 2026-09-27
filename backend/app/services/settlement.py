"""Module D1: bid settlement and margin simulation. See
docs/module-d1-plan.md for the full design -- schema rationale, the
per-line formula, the rates-are-authoritative rounding rule (§4), the
status lifecycle (§5a), segregation of duties (§5b, enforced generically in
app/services/approvals.py::decide -- not repeated here), the quantity check
at submit (§5c), and the worked example (§5).
"""

from __future__ import annotations

import hashlib
import io
from datetime import date, datetime, timezone
from decimal import ROUND_HALF_UP, Decimal
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.context import RequestContext
from app.core.enums import (
    ApprovalEntityType,
    AuditAction,
    BidSettlementStatus,
    QuotationLineItemStatus,
    SettlementCostSource,
)
from app.core.errors import ConflictError, ForbiddenError, NotFoundError, ValidationAppError
from app.boq.original_export import RejectedOriginalError, write_settled_rates
from app.integrations.s3 import get_object_bytes
from app.models.boq import BoqImportBatch, BoqLineItem
from app.models.quotation_ingestion import Quotation, QuotationLineItem
from app.models.settlement import (
    BidSettlement,
    BidSettlementLineItem,
    BidSettlementScenario,
    BidSettlementTradeOverride,
    SettlementReasonCode,
)
from app.models.vendors import Vendor
from app.schemas.approvals import ApprovalRequestCreate
from app.schemas.settlement import (
    ExportRequest,
    FxRateSet,
    LineCostUpdate,
    OutcomeRequest,
    QuantityMismatch,
    ScenarioCreate,
    SettlementDefaultsUpdate,
    SimulateLineOverride,
    SimulateLineResult,
    SimulateRequest,
    SimulateResult,
    TradeOverrideUpdate,
)
from app.services import approvals as approvals_service
from app.services import audit
from app.services import costlib as costlib_service

_SUBMIT_AND_OUTCOME_STATUSES = (BidSettlementStatus.SUBMITTED.value, BidSettlementStatus.APPROVED.value)

# lead_estimator+: structural decisions on the settlement itself (matches
# quotations' own WRITE_ROLES, migration 0016). estimator+ additionally
# covers day-to-day line costing/simulation. bd_director+ submits (an
# executive action); the routed approval tier (bd_director or
# managing_director) decides -- enforced generically by approvals.decide(),
# including segregation of duties, not re-checked here.
_HEADER_ROLES = ("lead_estimator", "procurement_head", "bd_director", "managing_director")
_LINE_ROLES = ("estimator", *_HEADER_ROLES)
_SUBMIT_ROLES = ("bd_director", "managing_director")

_Q2 = Decimal("0.01")
_Q3 = Decimal("0.001")


def _q2(x: Decimal) -> Decimal:
    return x.quantize(_Q2, rounding=ROUND_HALF_UP)


def _pct_fraction(pct: Decimal) -> Decimal:
    """Stored/input percentages are percent-numbers (10.00 means 10%); the
    formula needs the fraction."""
    return pct / Decimal(100)


def _require_role(ctx: RequestContext, roles: tuple[str, ...], action: str) -> None:
    if not ctx.has_role(*roles):
        raise ForbiddenError(f"{action} requires one of roles: {', '.join(roles)}")


async def _get_settlement(session: AsyncSession, settlement_id: UUID) -> BidSettlement:
    result = await session.execute(select(BidSettlement).where(BidSettlement.id == settlement_id))
    settlement = result.scalar_one_or_none()
    if settlement is None:
        raise NotFoundError(f"Bid settlement {settlement_id} not found")
    return settlement


def _require_draft(settlement: BidSettlement) -> None:
    if settlement.status != BidSettlementStatus.DRAFT.value:
        raise ConflictError(f"Bid settlement is {settlement.status}, not draft -- it can no longer be edited")


async def _get_line(session: AsyncSession, settlement_id: UUID, line_id: UUID) -> BidSettlementLineItem:
    result = await session.execute(
        select(BidSettlementLineItem).where(
            BidSettlementLineItem.id == line_id, BidSettlementLineItem.settlement_id == settlement_id
        )
    )
    line = result.scalar_one_or_none()
    if line is None:
        raise NotFoundError(f"Settlement line {line_id} not found")
    return line


async def list_settlement_lines(session: AsyncSession, settlement_id: UUID) -> list[BidSettlementLineItem]:
    result = await session.execute(
        select(BidSettlementLineItem).where(BidSettlementLineItem.settlement_id == settlement_id)
    )
    return list(result.scalars().all())


async def list_trade_overrides(session: AsyncSession, settlement_id: UUID) -> list[BidSettlementTradeOverride]:
    result = await session.execute(
        select(BidSettlementTradeOverride).where(BidSettlementTradeOverride.settlement_id == settlement_id)
    )
    return list(result.scalars().all())


async def _boq_items_by_id(session: AsyncSession, boq_line_item_ids: list[UUID]) -> dict[UUID, BoqLineItem]:
    if not boq_line_item_ids:
        return {}
    result = await session.execute(select(BoqLineItem).where(BoqLineItem.id.in_(boq_line_item_ids)))
    return {item.id: item for item in result.scalars().all()}


async def list_settlements(session: AsyncSession, project_id: UUID) -> list[BidSettlement]:
    result = await session.execute(
        select(BidSettlement).where(BidSettlement.project_id == project_id).order_by(BidSettlement.version_no)
    )
    return list(result.scalars().all())


async def get_settlement_detail(
    session: AsyncSession, settlement_id: UUID
) -> tuple[BidSettlement, list[BidSettlementLineItem], list[BidSettlementTradeOverride]]:
    settlement = await _get_settlement(session, settlement_id)
    lines = await list_settlement_lines(session, settlement_id)
    trade_overrides = await list_trade_overrides(session, settlement_id)
    return settlement, lines, trade_overrides


# --------------------------------------------------------------------------
# build draft
# --------------------------------------------------------------------------


async def build_settlement_draft(session: AsyncSession, ctx: RequestContext, project_id: UUID) -> BidSettlement:
    """New version for this project: auto-resolves each leaf BOQ line item
    (boq_quantity IS NOT NULL) from whichever QuotationLineItem is uniquely
    `accepted` for it; leaves it unresolved if zero or more than one vendor
    is accepted (see docs/module-d1-plan.md §2b -- a settlement never
    silently averages/auto-picks between competing accepted bids)."""
    _require_role(ctx, _HEADER_ROLES, "Building a settlement draft")

    max_version_no = (
        await session.execute(select(func.max(BidSettlement.version_no)).where(BidSettlement.project_id == project_id))
    ).scalar_one()
    version_no = (max_version_no or 0) + 1

    existing_current = (
        await session.execute(
            select(BidSettlement).where(BidSettlement.project_id == project_id, BidSettlement.is_current.is_(True))
        )
    ).scalar_one_or_none()
    if existing_current is not None:
        existing_current.is_current = False
        await session.flush()

    settlement = BidSettlement(
        tenant_id=ctx.tenant_id, project_id=project_id, version_no=version_no, is_current=True,
        status=BidSettlementStatus.DRAFT.value,
    )
    session.add(settlement)
    await session.flush()

    items_result = await session.execute(
        select(BoqLineItem).where(BoqLineItem.project_id == project_id, BoqLineItem.boq_quantity.is_not(None))
    )
    items = list(items_result.scalars().all())
    item_ids = [item.id for item in items]

    accepted_by_boq_item: dict[UUID, list[tuple[QuotationLineItem, Quotation]]] = {}
    if item_ids:
        accepted_result = await session.execute(
            select(QuotationLineItem, Quotation)
            .join(Quotation, Quotation.id == QuotationLineItem.quotation_id)
            .where(
                Quotation.project_id == project_id,
                QuotationLineItem.status == QuotationLineItemStatus.ACCEPTED.value,
                QuotationLineItem.boq_line_item_id.in_(item_ids),
            )
        )
        for line_item, quotation in accepted_result.all():
            accepted_by_boq_item.setdefault(line_item.boq_line_item_id, []).append((line_item, quotation))

    for item in items:
        candidates = accepted_by_boq_item.get(item.id, [])
        line = BidSettlementLineItem(
            tenant_id=ctx.tenant_id, settlement_id=settlement.id, boq_line_item_id=item.id, project_id=project_id,
            quantity=item.boq_quantity, quantity_at_build=item.boq_quantity,
        )
        if len(candidates) == 1:
            line_item, quotation = candidates[0]
            line.direct_unit_cost = line_item.unit_price
            line.source_currency = quotation.currency or settlement.currency
            line.cost_source = SettlementCostSource.QUOTATION_LINE.value
            line.source_quotation_line_item_id = line_item.id
            line.source_set_by = ctx.user_id
            line.source_set_at = datetime.now(timezone.utc)
        # 0 candidates: left unresolved for manual/cost-library entry.
        # >1 candidates: several vendors accepted for the same BOQ item --
        # left unresolved deliberately; a reviewer must call set_line_cost
        # to pick one explicitly.
        session.add(line)
    await session.flush()

    await audit.record(
        session, ctx, action=AuditAction.CREATE, entity_type="bid_settlement", entity_id=settlement.id,
        project_id=project_id, payload={"version_no": version_no, "line_count": len(items)},
    )
    return settlement


# --------------------------------------------------------------------------
# defaults / trade overrides / line cost / fx
# --------------------------------------------------------------------------


async def set_defaults(
    session: AsyncSession, ctx: RequestContext, settlement_id: UUID, data: SettlementDefaultsUpdate
) -> BidSettlement:
    _require_role(ctx, _HEADER_ROLES, "Updating settlement defaults")
    settlement = await _get_settlement(session, settlement_id)
    _require_draft(settlement)
    changes = data.model_dump(exclude_unset=True)
    for key, value in changes.items():
        setattr(settlement, key, value)
    await session.flush()
    await audit.record(
        session, ctx, action=AuditAction.UPDATE, entity_type="bid_settlement", entity_id=settlement.id,
        project_id=settlement.project_id, payload=changes,
    )
    return settlement


async def set_trade_override(
    session: AsyncSession, ctx: RequestContext, settlement_id: UUID, trade_node_id: UUID, data: TradeOverrideUpdate
) -> BidSettlementTradeOverride:
    _require_role(ctx, _HEADER_ROLES, "Setting a trade override")
    settlement = await _get_settlement(session, settlement_id)
    _require_draft(settlement)
    changes = data.model_dump(exclude_unset=True)

    result = await session.execute(
        select(BidSettlementTradeOverride).where(
            BidSettlementTradeOverride.settlement_id == settlement_id,
            BidSettlementTradeOverride.trade_node_id == trade_node_id,
        )
    )
    row = result.scalar_one_or_none()
    if row is None:
        row = BidSettlementTradeOverride(
            tenant_id=ctx.tenant_id, settlement_id=settlement_id, project_id=settlement.project_id,
            trade_node_id=trade_node_id,
        )
        session.add(row)
    for key, value in changes.items():
        setattr(row, key, value)
    await session.flush()
    await audit.record(
        session, ctx, action=AuditAction.UPDATE, entity_type="bid_settlement_trade_override", entity_id=row.id,
        project_id=settlement.project_id, payload={"trade_node_id": str(trade_node_id), **changes},
    )
    return row


async def set_line_cost(
    session: AsyncSession, ctx: RequestContext, settlement_id: UUID, line_id: UUID, data: LineCostUpdate
) -> BidSettlementLineItem:
    _require_role(ctx, _LINE_ROLES, "Setting a settlement line's cost")
    settlement = await _get_settlement(session, settlement_id)
    _require_draft(settlement)
    line = await _get_line(session, settlement_id, line_id)

    if data.cost_source is not None:
        if data.cost_source == SettlementCostSource.QUOTATION_LINE.value:
            if data.source_quotation_line_item_id is None:
                raise ValidationAppError("source_quotation_line_item_id is required for cost_source=quotation_line")
            row = (
                await session.execute(
                    select(QuotationLineItem, Quotation)
                    .join(Quotation, Quotation.id == QuotationLineItem.quotation_id)
                    .where(QuotationLineItem.id == data.source_quotation_line_item_id)
                )
            ).first()
            if row is None:
                raise NotFoundError(f"Quotation line item {data.source_quotation_line_item_id} not found")
            candidate, quotation = row
            if candidate.status != QuotationLineItemStatus.ACCEPTED.value:
                raise ValidationAppError("Only an accepted quotation line item can be used as a settlement line's cost source")
            if candidate.boq_line_item_id != line.boq_line_item_id:
                raise ValidationAppError("That quotation line item is not for this settlement line's BOQ item")
            line.direct_unit_cost = candidate.unit_price
            line.source_currency = quotation.currency or settlement.currency
            line.cost_source = SettlementCostSource.QUOTATION_LINE.value
            line.source_quotation_line_item_id = candidate.id
            line.source_cost_item_rate_id = None
            line.source_rate_as_of_date = None
        elif data.cost_source == SettlementCostSource.COST_LIBRARY_RATE.value:
            if data.cost_item_id is None or data.valid_on is None:
                raise ValidationAppError("cost_item_id and valid_on are required for cost_source=cost_library_rate")
            rate = await costlib_service.get_rate_as_of(session, ctx, data.cost_item_id, valid_on=data.valid_on)
            if rate is None:
                raise NotFoundError(f"No cost-library rate for item {data.cost_item_id} as of {data.valid_on}")
            line.direct_unit_cost = rate.total_rate
            line.source_currency = rate.currency
            line.cost_source = SettlementCostSource.COST_LIBRARY_RATE.value
            line.source_cost_item_rate_id = rate.id
            line.source_rate_as_of_date = data.valid_on
            line.source_quotation_line_item_id = None
        elif data.cost_source == SettlementCostSource.MANUAL.value:
            if data.manual_unit_cost is None or not data.source_note:
                raise ValidationAppError("manual_unit_cost and source_note (why) are required for cost_source=manual")
            line.direct_unit_cost = data.manual_unit_cost
            line.source_currency = data.manual_currency or settlement.currency
            line.cost_source = SettlementCostSource.MANUAL.value
            line.source_quotation_line_item_id = None
            line.source_cost_item_rate_id = None
            line.source_rate_as_of_date = None
        line.source_set_by = ctx.user_id
        line.source_set_at = datetime.now(timezone.utc)
        if data.source_note is not None:
            line.source_note = data.source_note

    pct_changes = data.model_dump(
        exclude_unset=True,
        include={"plant_pct_override", "overhead_pct_override", "volatility_pct_override", "markup_pct_override", "line_note"},
    )
    for key, value in pct_changes.items():
        setattr(line, key, value)

    await session.flush()
    await audit.record(
        session, ctx, action=AuditAction.UPDATE, entity_type="bid_settlement_line_item", entity_id=line.id,
        project_id=settlement.project_id,
        payload={"cost_source": line.cost_source, "direct_unit_cost": str(line.direct_unit_cost) if line.direct_unit_cost is not None else None},
    )
    return line


async def set_line_fx_rate(
    session: AsyncSession, ctx: RequestContext, settlement_id: UUID, line_id: UUID, data: FxRateSet
) -> BidSettlementLineItem:
    _require_role(ctx, _HEADER_ROLES, "Recording an FX rate")
    settlement = await _get_settlement(session, settlement_id)
    _require_draft(settlement)
    line = await _get_line(session, settlement_id, line_id)
    line.fx_rate = data.fx_rate
    line.fx_rate_date = data.fx_rate_date
    line.fx_recorded_by = ctx.user_id
    await session.flush()
    await audit.record(
        session, ctx, action=AuditAction.UPDATE, entity_type="bid_settlement_line_item", entity_id=line.id,
        project_id=settlement.project_id, payload={"fx_rate": str(data.fx_rate), "fx_rate_date": data.fx_rate_date.isoformat()},
    )
    return line


# --------------------------------------------------------------------------
# quantity check / refresh (§5c)
# --------------------------------------------------------------------------


async def _find_quantity_mismatches(session: AsyncSession, settlement: BidSettlement) -> list[QuantityMismatch]:
    lines = await list_settlement_lines(session, settlement.id)
    items = await _boq_items_by_id(session, [line.boq_line_item_id for line in lines])
    mismatches: list[QuantityMismatch] = []
    for line in lines:
        current_qty = items[line.boq_line_item_id].boq_quantity
        if current_qty is not None and Decimal(str(current_qty)) != Decimal(str(line.quantity)):
            mismatches.append(
                QuantityMismatch(
                    boq_line_item_id=line.boq_line_item_id, settlement_quantity=line.quantity, current_boq_quantity=current_qty
                )
            )
    return mismatches


async def refresh_quantities(session: AsyncSession, ctx: RequestContext, settlement_id: UUID) -> list[BidSettlementLineItem]:
    """Sets quantity to the current BOQ value on every stale line -- and
    touches only quantity: direct_unit_cost, cost_source, and every
    percentage override are left exactly as they were (§5c)."""
    _require_role(ctx, _HEADER_ROLES, "Refreshing settlement quantities")
    settlement = await _get_settlement(session, settlement_id)
    _require_draft(settlement)
    lines = await list_settlement_lines(session, settlement_id)
    items = await _boq_items_by_id(session, [line.boq_line_item_id for line in lines])

    changed = []
    for line in lines:
        current_qty = items[line.boq_line_item_id].boq_quantity
        if current_qty is not None and Decimal(str(current_qty)) != Decimal(str(line.quantity)):
            changed.append((line.boq_line_item_id, line.quantity, current_qty))
            line.quantity = current_qty
    settlement.quantities_refreshed_at = datetime.now(timezone.utc)
    await session.flush()
    await audit.record(
        session, ctx, action=AuditAction.UPDATE, entity_type="bid_settlement", entity_id=settlement.id,
        project_id=settlement.project_id,
        payload={
            "refreshed_lines": [
                {"boq_line_item_id": str(b), "old_quantity": str(o), "new_quantity": str(n)} for b, o, n in changed
            ]
        },
    )
    return lines


# --------------------------------------------------------------------------
# percentage resolution (§3) and the per-line formula (§4)
# --------------------------------------------------------------------------


def _resolve_pct(
    pct_type: str, *, settlement: BidSettlement, item: BoqLineItem, line: BidSettlementLineItem,
    trade_override_rows: dict[UUID, BidSettlementTradeOverride], data: SimulateRequest | None,
) -> Decimal:
    """Line override > trade override > project default. `data`, when
    given (simulate()'s transient what-if inputs), replaces the persisted
    value at whichever tier it targets -- a tier untouched by `data` still
    falls through to its persisted value, so simulate() with an empty body
    reproduces the settlement's own persisted resolution exactly."""
    persisted_line_val = getattr(line, f"{pct_type}_pct_override")
    transient_line = data.line_overrides.get(line.boq_line_item_id) if data else None
    transient_line_val = getattr(transient_line, f"{pct_type}_pct") if transient_line else None
    line_val = transient_line_val if transient_line_val is not None else persisted_line_val
    if line_val is not None:
        return Decimal(str(line_val))

    trade_val = None
    if item.trade_node_id is not None:
        transient_trade = data.trade_overrides.get(item.trade_node_id) if data else None
        transient_trade_val = getattr(transient_trade, f"{pct_type}_pct") if transient_trade else None
        row = trade_override_rows.get(item.trade_node_id)
        persisted_trade_val = getattr(row, f"{pct_type}_pct") if row else None
        trade_val = transient_trade_val if transient_trade_val is not None else persisted_trade_val
    if trade_val is not None:
        return Decimal(str(trade_val))

    transient_default = getattr(data, f"default_{pct_type}_pct") if data else None
    default_val = getattr(settlement, f"default_{pct_type}_pct")
    return Decimal(str(transient_default)) if transient_default is not None else Decimal(str(default_val))


def _calc_line(
    direct_unit_cost: Decimal, quantity: Decimal, plant_pct: Decimal, overhead_pct: Decimal,
    volatility_pct: Decimal, markup_pct: Decimal,
) -> tuple[Decimal, Decimal, Decimal, Decimal, Decimal, Decimal]:
    """§4's exact, full-precision model -- returns (base, plant, overhead,
    volatility, markup, model_sell). Never rounded here."""
    base = direct_unit_cost * quantity
    plant = base * _pct_fraction(plant_pct)
    overhead = base * _pct_fraction(overhead_pct)
    subtotal_1 = base + plant + overhead
    volatility = subtotal_1 * _pct_fraction(volatility_pct)
    subtotal_2 = subtotal_1 + volatility
    markup = subtotal_2 * _pct_fraction(markup_pct)
    model_sell = subtotal_2 + markup
    return base, plant, overhead, volatility, markup, model_sell


async def _compute_breakdown(
    session: AsyncSession, settlement: BidSettlement, data: SimulateRequest | None
) -> SimulateResult:
    lines = await list_settlement_lines(session, settlement.id)
    trade_override_rows = {row.trade_node_id: row for row in await list_trade_overrides(session, settlement.id)}
    items = await _boq_items_by_id(session, [line.boq_line_item_id for line in lines])

    line_results: list[SimulateLineResult] = []
    unresolved: list[UUID] = []
    direct_cost_total = plant_total = overhead_total = volatility_total = markup_total = Decimal(0)
    exact_model_total = tender_total_raw = Decimal(0)

    for line in lines:
        item = items[line.boq_line_item_id]
        plant_pct = _resolve_pct("plant", settlement=settlement, item=item, line=line, trade_override_rows=trade_override_rows, data=data)
        overhead_pct = _resolve_pct("overhead", settlement=settlement, item=item, line=line, trade_override_rows=trade_override_rows, data=data)
        volatility_pct = _resolve_pct("volatility", settlement=settlement, item=item, line=line, trade_override_rows=trade_override_rows, data=data)
        markup_pct = _resolve_pct("markup", settlement=settlement, item=item, line=line, trade_override_rows=trade_override_rows, data=data)

        if line.direct_unit_cost is None:
            unresolved.append(line.boq_line_item_id)
            line_results.append(
                SimulateLineResult(
                    boq_line_item_id=line.boq_line_item_id, resolved=False, quantity=Decimal(str(line.quantity)),
                    direct_unit_cost=None, plant_pct=plant_pct, overhead_pct=overhead_pct, volatility_pct=volatility_pct,
                    markup_pct=markup_pct, base=Decimal(0), plant=Decimal(0), overhead=Decimal(0), volatility=Decimal(0),
                    markup=Decimal(0), model_sell=Decimal(0), unit_sell_rate=None, line_amount=None,
                )
            )
            continue

        quantity = Decimal(str(line.quantity))
        direct_unit_cost = Decimal(str(line.direct_unit_cost))
        base, plant, overhead, volatility, markup, model_sell = _calc_line(
            direct_unit_cost, quantity, plant_pct, overhead_pct, volatility_pct, markup_pct
        )
        unit_sell_rate = _q2(model_sell / quantity) if quantity != 0 else Decimal("0.00")
        line_amount = _q2(unit_sell_rate * quantity)

        direct_cost_total += base
        plant_total += plant
        overhead_total += overhead
        volatility_total += volatility
        markup_total += markup
        exact_model_total += model_sell
        tender_total_raw += line_amount

        line_results.append(
            SimulateLineResult(
                boq_line_item_id=line.boq_line_item_id, resolved=True, quantity=quantity, direct_unit_cost=direct_unit_cost,
                plant_pct=plant_pct, overhead_pct=overhead_pct, volatility_pct=volatility_pct, markup_pct=markup_pct,
                base=base, plant=plant, overhead=overhead, volatility=volatility, markup=markup, model_sell=model_sell,
                unit_sell_rate=unit_sell_rate, line_amount=line_amount,
            )
        )

    tender_total = _q2(tender_total_raw)
    rounding_difference = _q2(tender_total - exact_model_total)
    margin_on_sell_pct = (markup_total / tender_total * 100).quantize(_Q3, rounding=ROUND_HALF_UP) if tender_total != 0 else None

    required_role = await approvals_service.preview_required_role(
        session, ApprovalEntityType.BID_SUBMISSION.value, tender_total,
        margin_pct=margin_on_sell_pct,
    )

    return SimulateResult(
        lines=line_results, unresolved_line_ids=unresolved, direct_cost_total=_q2(direct_cost_total),
        plant_total=_q2(plant_total), overhead_total=_q2(overhead_total), volatility_total=_q2(volatility_total),
        markup_total=_q2(markup_total), exact_model_total=_q2(exact_model_total), tender_total=tender_total,
        rounding_difference=rounding_difference, margin_on_sell_pct=margin_on_sell_pct, required_role=required_role,
    )


async def simulate(session: AsyncSession, ctx: RequestContext, settlement_id: UUID, data: SimulateRequest) -> SimulateResult:
    """Stateless -- no DB writes. See docs/module-d1-plan.md §1/§4."""
    _require_role(ctx, _LINE_ROLES, "Simulating a settlement")
    settlement = await _get_settlement(session, settlement_id)
    return await _compute_breakdown(session, settlement, data)


async def save_scenario(
    session: AsyncSession, ctx: RequestContext, settlement_id: UUID, data: ScenarioCreate
) -> BidSettlementScenario:
    _require_role(ctx, _LINE_ROLES, "Saving a settlement scenario")
    settlement = await _get_settlement(session, settlement_id)
    result = await _compute_breakdown(session, settlement, data.inputs)
    scenario = BidSettlementScenario(
        tenant_id=ctx.tenant_id, settlement_id=settlement.id, project_id=settlement.project_id, label=data.label,
        inputs=data.inputs.model_dump(mode="json"), result=result.model_dump(mode="json"), created_by=ctx.user_id,
    )
    session.add(scenario)
    await session.flush()
    return scenario


async def list_scenarios(session: AsyncSession, settlement_id: UUID) -> list[BidSettlementScenario]:
    result = await session.execute(
        select(BidSettlementScenario).where(BidSettlementScenario.settlement_id == settlement_id).order_by(BidSettlementScenario.created_at)
    )
    return list(result.scalars().all())


# --------------------------------------------------------------------------
# submit / decide (§5a, §5c)
# --------------------------------------------------------------------------


async def submit_settlement(session: AsyncSession, ctx: RequestContext, settlement_id: UUID) -> BidSettlement:
    _require_role(ctx, _SUBMIT_ROLES, "Submitting a settlement")
    settlement = await _get_settlement(session, settlement_id)
    _require_draft(settlement)

    mismatches = await _find_quantity_mismatches(session, settlement)
    if mismatches:
        raise ConflictError(
            "Settlement quantities are stale against the current BOQ -- call refresh-quantities first",
            mismatches=[m.model_dump(mode="json") for m in mismatches],
        )

    lines = await list_settlement_lines(session, settlement.id)
    unresolved = [line.boq_line_item_id for line in lines if line.direct_unit_cost is None]
    if unresolved:
        raise ConflictError(
            "Every settlement line needs a resolved cost before submission",
            unresolved_line_ids=[str(i) for i in unresolved],
        )

    fx_missing = [
        line.boq_line_item_id for line in lines if line.source_currency != settlement.currency and line.fx_rate is None
    ]
    if fx_missing:
        raise ConflictError(
            "Every currency-mismatched line needs a recorded FX rate before submission",
            fx_missing_line_ids=[str(i) for i in fx_missing],
        )

    breakdown = await _compute_breakdown(session, settlement, data=None)
    result_by_item = {r.boq_line_item_id: r for r in breakdown.lines}
    for line in lines:
        r = result_by_item[line.boq_line_item_id]
        line.unit_sell_rate = r.unit_sell_rate
        line.line_amount = r.line_amount

    settlement.direct_cost_total = breakdown.direct_cost_total
    settlement.plant_total = breakdown.plant_total
    settlement.overhead_total = breakdown.overhead_total
    settlement.volatility_total = breakdown.volatility_total
    settlement.markup_total = breakdown.markup_total
    settlement.tender_total = breakdown.tender_total
    settlement.rounding_difference = breakdown.rounding_difference
    settlement.margin_on_sell_pct = breakdown.margin_on_sell_pct
    settlement.status = BidSettlementStatus.SUBMITTED.value
    settlement.submitted_at = datetime.now(timezone.utc)
    settlement.submitted_by = ctx.user_id
    await session.flush()

    request = await approvals_service.create_request(
        session, ctx,
        ApprovalRequestCreate(
            project_id=settlement.project_id, entity_type=ApprovalEntityType.BID_SUBMISSION.value,
            entity_id=settlement.id, amount=breakdown.tender_total, currency=settlement.currency,
            margin_pct=breakdown.margin_on_sell_pct,
            payload_snapshot={"version_no": settlement.version_no, "tender_total": str(breakdown.tender_total)},
        ),
    )
    settlement.approval_request_id = request.id
    await session.flush()

    await audit.record(
        session, ctx, action=AuditAction.UPDATE, entity_type="bid_settlement", entity_id=settlement.id,
        project_id=settlement.project_id,
        payload={"status": settlement.status, "tender_total": str(settlement.tender_total), "approval_request_id": str(request.id)},
    )
    return settlement


async def decide_settlement(
    session: AsyncSession, ctx: RequestContext, settlement_id: UUID, *, approve: bool, note: str | None
) -> BidSettlement:
    """Thin wrapper over approvals.decide() -- role check, segregation of
    duties, and MFA step-up are all enforced there, generically, not
    repeated here (§5b). This just syncs bid_settlements.status/decided_*
    and audits the settlement-level transition in the same call."""
    settlement = await _get_settlement(session, settlement_id)
    if settlement.status != BidSettlementStatus.SUBMITTED.value:
        raise ConflictError(f"Bid settlement is {settlement.status}, not submitted -- nothing to decide")
    if settlement.approval_request_id is None:
        raise ConflictError("This settlement has no approval request to decide")

    request = await approvals_service.decide(session, ctx, settlement.approval_request_id, approve=approve, note=note)

    settlement.decided_at = datetime.now(timezone.utc)
    settlement.decided_by = ctx.user_id
    if request.status == "approved":
        settlement.status = BidSettlementStatus.APPROVED.value
    elif request.status == "rejected":
        settlement.status = BidSettlementStatus.REJECTED.value
    await session.flush()

    await audit.record(
        session, ctx, action=AuditAction.APPROVE if approve else AuditAction.REJECT, entity_type="bid_settlement",
        entity_id=settlement.id, project_id=settlement.project_id, payload={"status": settlement.status, "note": note},
    )
    return settlement


# --------------------------------------------------------------------------
# Module D2: generated export (§3 of docs/module-d2-plan.md)
# --------------------------------------------------------------------------


def _build_export_workbook(
    settlement: BidSettlement, all_items: list[BoqLineItem], lines_by_item_id: dict[UUID, BidSettlementLineItem],
    data: ExportRequest,
) -> bytes:
    """A fresh workbook, built entirely from the six columns below -- never
    opens or touches any retained original (that's D3). No document
    properties, comments, defined names, or extra sheets are ever set, so
    there is nothing beyond these cells for the leak test (§4 of the plan)
    to have to find clean."""
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    header = ("Item No", "Description", "Unit", "Qty", "Rate", "Amount")
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    row = 1
    for item in all_items:
        row += 1
        line = lines_by_item_id.get(item.id)
        ws.cell(row=row, column=1, value=item.item_no)
        ws.cell(row=row, column=2, value=item.description)
        ws.cell(row=row, column=3, value=item.uom)
        qty = line.quantity if line is not None else item.boq_quantity
        ws.cell(row=row, column=4, value=float(qty) if qty is not None else None)
        if line is not None and line.unit_sell_rate is not None:
            ws.cell(row=row, column=5, value=float(line.unit_sell_rate))
            ws.cell(row=row, column=6, value=float(line.line_amount))

    last_data_row = row
    subtotal_row = last_data_row + 2
    ws.cell(row=subtotal_row, column=2, value="Subtotal").font = Font(bold=True)
    ws.cell(row=subtotal_row, column=6, value=f"=SUM(F2:F{last_data_row})")

    if data.include_vat:
        vat_row = subtotal_row + 1
        ws.cell(row=vat_row, column=2, value=f"VAT @ {data.vat_pct:g}%")
        ws.cell(row=vat_row, column=6, value=f"=F{subtotal_row}*{data.vat_pct}/100")
        total_row = vat_row + 1
        ws.cell(row=total_row, column=2, value="Total incl. VAT").font = Font(bold=True)
        ws.cell(row=total_row, column=6, value=f"=F{subtotal_row}+F{vat_row}").font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def export_settlement(
    session: AsyncSession, ctx: RequestContext, settlement_id: UUID, data: ExportRequest
) -> tuple[bytes, str, str]:
    """Returns (file_bytes, sha256, filename). Export only from an
    approved settlement; audited with the sha256 (§3)."""
    _require_role(ctx, _LINE_ROLES, "Exporting a settlement")
    settlement = await _get_settlement(session, settlement_id)
    if settlement.status != BidSettlementStatus.APPROVED.value:
        raise ConflictError(f"Bid settlement is {settlement.status}, not approved -- export is only available once approved")

    lines = await list_settlement_lines(session, settlement.id)
    lines_by_item_id = {line.boq_line_item_id: line for line in lines}
    all_items = list(
        (
            await session.execute(
                select(BoqLineItem).where(BoqLineItem.project_id == settlement.project_id).order_by(BoqLineItem.path)
            )
        ).scalars().all()
    )

    file_bytes = _build_export_workbook(settlement, all_items, lines_by_item_id, data)
    sha256 = hashlib.sha256(file_bytes).hexdigest()
    filename = f"settlement-{settlement.project_id}-v{settlement.version_no}.xlsx"

    await audit.record(
        session, ctx, action=AuditAction.EXPORT, entity_type="bid_settlement", entity_id=settlement.id,
        project_id=settlement.project_id,
        payload={
            "version_no": settlement.version_no, "sha256": sha256, "include_vat": data.include_vat,
            "vat_pct": data.vat_pct if data.include_vat else None,
        },
    )
    return file_bytes, sha256, filename


# --------------------------------------------------------------------------
# Module D2: win/loss (§5 of docs/module-d2-plan.md)
# --------------------------------------------------------------------------

_OUTCOME_ROLES = ("bd_director", "managing_director")


async def list_reason_codes(
    session: AsyncSession, ctx: RequestContext, *, include_inactive: bool = False
) -> list[SettlementReasonCode]:
    stmt = select(SettlementReasonCode).where(SettlementReasonCode.tenant_id == ctx.tenant_id)
    if not include_inactive:
        stmt = stmt.where(SettlementReasonCode.is_active.is_(True))
    result = await session.execute(stmt.order_by(SettlementReasonCode.code))
    return list(result.scalars().all())


async def create_reason_code(session: AsyncSession, ctx: RequestContext, code: str, label: str) -> SettlementReasonCode:
    _require_role(ctx, _HEADER_ROLES, "Creating a settlement reason code")
    row = SettlementReasonCode(tenant_id=ctx.tenant_id, code=code, label=label)
    session.add(row)
    await session.flush()
    await audit.record(session, ctx, action=AuditAction.CREATE, entity_type="settlement_reason_code", entity_id=row.id, payload={"code": code})
    return row


async def update_reason_code(
    session: AsyncSession, ctx: RequestContext, reason_code_id: UUID, *, label: str | None = None, is_active: bool | None = None
) -> SettlementReasonCode:
    _require_role(ctx, _HEADER_ROLES, "Updating a settlement reason code")
    result = await session.execute(select(SettlementReasonCode).where(SettlementReasonCode.id == reason_code_id))
    row = result.scalar_one_or_none()
    if row is None:
        raise NotFoundError(f"Reason code {reason_code_id} not found")
    if label is not None:
        row.label = label
    if is_active is not None:
        row.is_active = is_active
    await session.flush()
    await audit.record(
        session, ctx, action=AuditAction.UPDATE, entity_type="settlement_reason_code", entity_id=row.id,
        payload={"label": label, "is_active": is_active},
    )
    return row


async def record_outcome(
    session: AsyncSession, ctx: RequestContext, settlement_id: UUID, data: OutcomeRequest
) -> BidSettlement:
    _require_role(ctx, _OUTCOME_ROLES, "Recording a settlement outcome")
    settlement = await _get_settlement(session, settlement_id)
    if settlement.status not in _SUBMIT_AND_OUTCOME_STATUSES:
        raise ConflictError(
            f"Bid settlement is {settlement.status} -- outcome can only be recorded from submitted or approved"
        )

    if data.reason_codes:
        valid = {
            row.code
            for row in (
                await session.execute(
                    select(SettlementReasonCode).where(
                        SettlementReasonCode.tenant_id == ctx.tenant_id, SettlementReasonCode.is_active.is_(True),
                        SettlementReasonCode.code.in_(data.reason_codes),
                    )
                )
            ).scalars().all()
        }
        unknown = set(data.reason_codes) - valid
        if unknown:
            raise ValidationAppError(f"Unknown or inactive reason code(s): {', '.join(sorted(unknown))}")

    if data.competitor_vendor_ids:
        found = {
            row.id
            for row in (
                await session.execute(select(Vendor).where(Vendor.id.in_(data.competitor_vendor_ids)))
            ).scalars().all()
        }
        missing = set(data.competitor_vendor_ids) - found
        if missing:
            raise ValidationAppError(f"Unknown vendor id(s): {', '.join(str(v) for v in sorted(missing, key=str))}")

    settlement.outcome = data.outcome
    settlement.status = data.outcome  # "won"/"lost" -- moving to won/lost sets the settlement status
    settlement.outcome_our_price = data.our_price if data.our_price is not None else settlement.tender_total
    settlement.outcome_winning_price = data.winning_price
    settlement.outcome_competitor_names = data.competitor_names
    settlement.outcome_competitor_vendor_ids = data.competitor_vendor_ids
    settlement.outcome_reason_codes = data.reason_codes
    settlement.outcome_recorded_by = ctx.user_id
    settlement.outcome_recorded_at = datetime.now(timezone.utc)
    settlement.outcome_note = data.note
    await session.flush()

    await audit.record(
        session, ctx, action=AuditAction.UPDATE, entity_type="bid_settlement", entity_id=settlement.id,
        project_id=settlement.project_id,
        payload={"outcome": data.outcome, "reason_codes": data.reason_codes, "our_price": str(settlement.outcome_our_price)},
    )
    return settlement


# --------------------------------------------------------------------------
# Module D3: export into the client's original workbook
# --------------------------------------------------------------------------

_LEAK_CHECK_FIELDS = ("direct_unit_cost", "cost_source", "fx_rate", "source_note")


def _leak_denylist(lines: list[BidSettlementLineItem]) -> list[str]:
    """Runtime safety net, not the authoritative check (that's the
    dedicated scanner in tests/unit/test_settlement_export_leak.py,
    exercised against every export path) -- D3 only ever writes two float
    values per line, so this should never actually fire; it exists in
    case a future change to write_settled_rates ever serializes more than
    that."""
    values: list[str] = []
    for line in lines:
        for field_name in _LEAK_CHECK_FIELDS:
            value = getattr(line, field_name)
            if value is not None and str(value).strip():
                values.append(str(value))
    return values


async def _get_single_eligible_batch(session: AsyncSession, settlement: BidSettlement, lines: list[BidSettlementLineItem]):
    """§3 of docs/module-d3-plan.md: every line must trace to the same
    single import batch, that batch must have retained the original
    workbook, and must know its rate column. Any failure is a clear,
    named 409 -- the generated export (D2) always remains available
    regardless."""
    boq_items = await _boq_items_by_id(session, [line.boq_line_item_id for line in lines])
    batch_ids = {boq_items[line.boq_line_item_id].import_batch_id for line in lines}
    if len(batch_ids) != 1 or None in batch_ids:
        raise ConflictError(
            "This settlement's BOQ lines don't all trace to a single import batch -- original-workbook export "
            "isn't available; use the generated export instead."
        )
    (batch_id,) = batch_ids
    batch = (await session.execute(select(BoqImportBatch).where(BoqImportBatch.id == batch_id))).scalar_one_or_none()
    if batch is None or batch.source_object_key is None:
        raise ConflictError(
            "This settlement's BOQ import didn't retain the original workbook -- use the generated export instead."
        )
    if batch.rate_column is None:
        raise ConflictError(
            "This settlement's BOQ import didn't record a rate column -- use the generated export instead."
        )
    return batch, boq_items


async def _run_original_export(
    session: AsyncSession, settlement: BidSettlement, lines: list[BidSettlementLineItem]
):
    """Shared by preview and the real export -- always a real write
    attempt (in memory), never partially implemented, so the preview's
    report is never optimistic about what the real export would do."""
    batch, boq_items = await _get_single_eligible_batch(session, settlement, lines)
    original_bytes = await get_object_bytes(batch.source_object_key)

    writes = [
        (boq_items[line.boq_line_item_id].source_row_number, float(line.unit_sell_rate), float(line.line_amount))
        for line in lines
        if boq_items[line.boq_line_item_id].import_batch_id == batch.id
        and boq_items[line.boq_line_item_id].source_row_number is not None
    ]
    try:
        output_bytes, report = write_settled_rates(
            original_bytes, sheet_name=batch.sheet_name, rate_column=batch.rate_column,
            amount_column=batch.amount_column, writes=writes,
        )
    except RejectedOriginalError as exc:
        raise ConflictError(str(exc)) from exc

    return batch, output_bytes, report


async def preview_original_export(session: AsyncSession, ctx: RequestContext, settlement_id: UUID) -> dict:
    _require_role(ctx, _LINE_ROLES, "Previewing an original-workbook export")
    settlement = await _get_settlement(session, settlement_id)
    if settlement.status != BidSettlementStatus.APPROVED.value:
        raise ConflictError(f"Bid settlement is {settlement.status}, not approved -- export is only available once approved")
    lines = await list_settlement_lines(session, settlement.id)

    _batch, _output_bytes, report = await _run_original_export(session, settlement, lines)
    leaked = [m for m in _leak_denylist(lines) if m.encode() in _output_bytes]
    if leaked:  # pragma: no cover -- defense in depth, see _leak_denylist's docstring
        raise ConflictError("Internal safety check failed: the generated file appears to contain non-price data.")
    return {"ok": report.ok, "lost_features": report.lost_features, "unexpected_cell_changes": report.unexpected_cell_changes}


async def export_original_settlement(
    session: AsyncSession, ctx: RequestContext, settlement_id: UUID, *, accept_loss: bool
) -> tuple[bytes, str, str]:
    _require_role(ctx, _LINE_ROLES, "Exporting a settlement into its original workbook")
    settlement = await _get_settlement(session, settlement_id)
    if settlement.status != BidSettlementStatus.APPROVED.value:
        raise ConflictError(f"Bid settlement is {settlement.status}, not approved -- export is only available once approved")
    lines = await list_settlement_lines(session, settlement.id)

    batch, output_bytes, report = await _run_original_export(session, settlement, lines)
    leaked = [m for m in _leak_denylist(lines) if m.encode() in output_bytes]
    if leaked:  # pragma: no cover -- defense in depth, see _leak_denylist's docstring
        raise ConflictError("Internal safety check failed: the generated file appears to contain non-price data.")

    if not report.ok and not accept_loss:
        raise ConflictError(
            "Writing into the original workbook would lose or change: " + ", ".join(report.lost_features) +
            " -- retry with accept_loss=true to proceed anyway, or use the generated export instead.",
            lost_features=report.lost_features,
        )

    sha256 = hashlib.sha256(output_bytes).hexdigest()
    filename = f"settlement-{settlement.project_id}-v{settlement.version_no}-original.xlsx"

    await audit.record(
        session, ctx, action=AuditAction.EXPORT, entity_type="bid_settlement", entity_id=settlement.id,
        project_id=settlement.project_id,
        payload={
            "version_no": settlement.version_no, "sha256": sha256, "export_kind": "original",
            "import_batch_id": str(batch.id), "fidelity_ok": report.ok, "lost_features": report.lost_features,
            "accepted_loss": accept_loss if not report.ok else None,
        },
    )
    return output_bytes, sha256, filename
