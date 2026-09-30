"""Verification report exports: xlsx (full detail, filterable) and PDF (summary + exceptions)."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any

from app.models.verification import VerificationLine, VerificationRun

CLASS_LABELS = {
    "match": "Match",
    "variance": "Variance",
    "missing_in_drawings": "Missing in drawings",
    "not_verifiable": "Not verifiable",
    "missing_in_boq": "Missing in BOQ",
}
METHOD_LABELS = {"measured": "Measured", "annotated": "Annotated", "none": "-"}
MAPPED_LABELS = {"ai": "AI-mapped", "rules": "Rules", "user": "Confirmed by estimator", "none": "-"}
PDF_ROW_LIMIT = 400


def _f(v: Any) -> float | None:
    return None if v is None else float(v)


def _evidence_text(line: VerificationLine) -> str:
    parts = []
    for e in line.evidence or []:
        where = f"p.{e['page']}" if e.get("page") else ""
        parts.append(f"{e.get('label', '')} = {e.get('value')} {e.get('unit') or ''} {where}".strip())
    return "; ".join(parts)


def build_xlsx(run: VerificationRun, lines: list[VerificationLine], project_label: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)
    ws["A1"], ws["A1"].font = f"Drawing vs BOQ verification: {project_label}", Font(bold=True, size=14)
    ws["A2"] = f"Generated {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC. Report id {run.id}."
    ws["A3"] = ("Quantities are measured from the drawings or read from their annotations. AI proposes which "
                "evidence belongs to a BOQ line; it never supplies a quantity.")
    row = 5
    ws.cell(row, 1, "Class").font = bold
    ws.cell(row, 2, "Lines").font = bold
    for key, label in CLASS_LABELS.items():
        row += 1
        ws.cell(row, 1, label)
        ws.cell(row, 2, (run.counts or {}).get(key, 0))
    row += 2
    ws.cell(row, 1, "Trade").font = bold
    for c, h in enumerate(("Lines", "Match", "Variance", "Missing in drawings", "Not verifiable"), start=2):
        ws.cell(row, c, h).font = bold
    for t in (run.summary or {}).get("by_trade", []):
        row += 1
        for c, k in enumerate(("trade", "lines", "match", "variance", "missing_in_drawings", "not_verifiable"), start=1):
            ws.cell(row, c, t.get(k))
    row += 2
    ws.cell(row, 1, "Warnings").font = bold
    for w in run.warnings or []:
        row += 1
        ws.cell(row, 1, w)
    ws.column_dimensions["A"].width = 44

    headers = ["Item", "Description", "Trade", "Class", "Unit", "BOQ qty", "Drawing qty", "Delta", "Delta %",
               "Confidence", "Method", "Mapping", "Evidence", "Note"]
    fills = {"variance": "FFF2CC", "missing_in_drawings": "F8CBAD", "missing_in_boq": "F8CBAD", "not_verifiable": "EDEDED", "match": "E2EFDA"}

    def sheet(title: str, rows: list[VerificationLine]) -> None:
        s = wb.create_sheet(title)
        s.append(headers)
        for c in range(1, len(headers) + 1):
            s.cell(1, c).font = bold
        for r in rows:
            s.append([
                r.item_no, r.description, r.trade_name, CLASS_LABELS.get(r.vclass, r.vclass), r.boq_uom, _f(r.boq_quantity),
                _f(r.drawing_quantity), _f(r.delta), _f(r.delta_pct), _f(r.confidence), METHOD_LABELS.get(r.method, r.method),
                MAPPED_LABELS.get(r.mapped_by, r.mapped_by), _evidence_text(r), r.note,
            ])
            s.cell(s.max_row, 4).fill = PatternFill("solid", fgColor=fills.get(r.vclass, "FFFFFF"))
            s.cell(s.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        widths = [10, 60, 18, 20, 8, 12, 12, 12, 10, 11, 12, 22, 60, 50]
        for i, w in enumerate(widths, start=1):
            s.column_dimensions[get_column_letter(i)].width = w
        s.freeze_panes = "C2"
        s.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2, s.max_row)}"

    sheet("Lines", [l for l in lines if l.boq_line_item_id is not None])
    sheet("Missing in BOQ", [l for l in lines if l.vclass == "missing_in_boq"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(run: VerificationRun, lines: list[VerificationLine], project_label: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small")
    small.fontSize = 7.5
    small.leading = 9
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24,
                            title=f"Verification: {project_label}")
    story: list[Any] = [
        Paragraph(f"Drawing vs BOQ verification", styles["Title"]),
        Paragraph(f"{project_label} &middot; generated {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC", styles["Normal"]),
        Spacer(1, 6),
        Paragraph("Quantities are measured from the drawings or read from their annotations. AI proposes which "
                  "evidence belongs to a BOQ line; it never supplies a quantity.", small),
        Spacer(1, 10),
    ]
    counts = run.counts or {}
    t = Table([["Class", "Lines"]] + [[CLASS_LABELS[k], counts.get(k, 0)] for k in CLASS_LABELS], hAlign="LEFT")
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDDDDD")), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                           ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    story += [t, Spacer(1, 10)]
    for w in run.warnings or []:
        story.append(Paragraph(f"Warning: {w}", small))
    trades = (run.summary or {}).get("by_trade", [])
    if trades:
        tt = Table([["Trade", "Lines", "Match", "Variance", "Missing in drawings", "Not verifiable"]] +
                   [[x["trade"], x["lines"], x["match"], x["variance"], x["missing_in_drawings"], x["not_verifiable"]] for x in trades],
                   hAlign="LEFT")
        tt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDDDDD")), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                                ("FONTSIZE", (0, 0), (-1, -1), 8)]))
        story += [Spacer(1, 10), Paragraph("By trade", styles["Heading3"]), tt]
    exceptions = [l for l in lines if l.vclass in ("variance", "missing_in_drawings", "missing_in_boq")]
    story += [Spacer(1, 12), Paragraph(f"Exceptions ({len(exceptions)})", styles["Heading3"])]
    data = [["Item", "Description", "Class", "Unit", "BOQ qty", "Drawing qty", "Delta %", "Conf.", "Mapping"]]
    for l in exceptions[:PDF_ROW_LIMIT]:
        data.append([
            l.item_no or "", Paragraph((l.description or "")[:220], small), CLASS_LABELS.get(l.vclass, l.vclass), l.boq_uom or "",
            "" if l.boq_quantity is None else f"{float(l.boq_quantity):,.2f}",
            "" if l.drawing_quantity is None else f"{float(l.drawing_quantity):,.2f}",
            "" if l.delta_pct is None else f"{float(l.delta_pct):+.1f}%", f"{float(l.confidence):.2f}",
            MAPPED_LABELS.get(l.mapped_by, l.mapped_by),
        ])
    et = Table(data, repeatRows=1, colWidths=[42, 300, 90, 34, 62, 62, 48, 36, 90])
    et.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDDDDD")), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                            ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(et)
    if len(exceptions) > PDF_ROW_LIMIT:
        story.append(Paragraph(f"Showing the first {PDF_ROW_LIMIT} of {len(exceptions)} exceptions. The xlsx export has every line.", small))
    doc.build(story)
    return buf.getvalue()
