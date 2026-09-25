"""Generates the .xlsx pricing sheet attached to every dispatched RFQ.

Two worksheets:
  - "Pricing" (visible): item_no/description/unit/quantity columns are
    filled in and cell-locked; rate/remarks columns are left blank and
    unlocked for the vendor to fill in. Row 2 also carries an unlocked
    currency cell (B2, pre-filled with the tenant's base currency) and an
    unlocked "prices include VAT?" cell (E2, left blank) -- both required
    before any line on the resulting quotation can be accepted (SRS change
    #4; see app/services/quotation_ingestion.py). Column G is a *hidden*,
    always-locked row-id column: G{row} holds that row's boq_line_item_id
    verbatim. Module C2's deterministic parser
    (app/procurement/pricing_sheet_parser.py) keys off this column, never
    off row position -- a vendor reordering, duplicating, deleting, or
    inserting rows is detected (the resulting row-id set won't match what
    we sent) and sends the whole quotation to needs-review instead of
    silently mis-mapping a rate to the wrong BOQ item (SRS change #1). Sheet
    protection is a UX guardrail (stops a vendor accidentally overtyping a
    quantity), not a security boundary -- it carries no password.
  - "_meta" (hidden): rfq_ref/rfq_id/reply_token (checked against the
    matched Rfq before any deterministic parse is trusted -- SRS change #1)
    plus the same row_index -> boq_line_item_id mapping as originally sent,
    kept only as a diagnostic record of the sheet's original layout, not as
    something the parser trusts for row identity.
"""

from __future__ import annotations

import hashlib
import io
import uuid

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.worksheet import Worksheet

from app.procurement.content import RfqLineItem

_HEADER_FILL = PatternFill(start_color="FFE0E0E0", end_color="FFE0E0E0", fill_type="solid")
_HEADER_FONT = Font(bold=True)
_LOCKED_COLUMNS = ("A", "B", "C", "D", "G")  # Item No, Description, Unit, Quantity, (hidden) row id
# Rate/Remarks (E, F) are the only data-row cells left unlocked -- see
# _write_pricing_sheet. Row 2's currency (B2) / VAT (E2) cells are also
# unlocked, as a row-2-specific override below.
FIRST_DATA_ROW = 4
_CURRENCY_ROW = 2


def _write_pricing_sheet(
    ws: Worksheet, package_name: str, rfq_ref: str, items: list[RfqLineItem], default_currency: str
) -> None:
    ws.title = "Pricing"
    ws["A1"] = f"RFQ {rfq_ref} - {package_name}"
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A1:F1")

    ws["A2"] = "Currency:"
    ws["A2"].font = _HEADER_FONT
    ws["B2"] = default_currency
    ws["D2"] = "Prices include VAT? (Yes/No)"
    ws["D2"].font = _HEADER_FONT
    ws["E2"] = None

    headers = ["Item No", "Description", "Unit", "Quantity", "Rate", "Remarks"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL

    for offset, item in enumerate(items):
        row = FIRST_DATA_ROW + offset
        ws.cell(row=row, column=1, value=item.item_no)
        ws.cell(row=row, column=2, value=item.description)
        ws.cell(row=row, column=3, value=item.uom or "")
        ws.cell(row=row, column=4, value=float(item.quantity) if item.quantity is not None else None)
        # Rate (E) and Remarks (F) are left blank for the vendor.
        ws.cell(row=row, column=7, value=str(item.boq_line_item_id))

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 30
    ws.column_dimensions["G"].hidden = True
    ws["B3"].alignment = Alignment(wrap_text=True)

    last_row = FIRST_DATA_ROW + max(len(items) - 1, 0)
    for row in range(1, last_row + 1):
        for col_letter in ("A", "B", "C", "D", "E", "F", "G"):
            cell = ws[f"{col_letter}{row}"]
            cell.protection = Protection(locked=col_letter in _LOCKED_COLUMNS or row < FIRST_DATA_ROW)
    # Row-2 override: currency/VAT are meant for the vendor to fill in, unlike
    # every other header-area cell (which the blanket "row < first data row"
    # rule above locks).
    ws[f"B{_CURRENCY_ROW}"].protection = Protection(locked=False)
    ws[f"E{_CURRENCY_ROW}"].protection = Protection(locked=False)
    ws.protection.sheet = True


def _write_meta_sheet(
    wb: Workbook, *, rfq_id: uuid.UUID, rfq_ref: str, reply_token: str, items: list[RfqLineItem]
) -> None:
    ws = wb.create_sheet("_meta")
    ws.append(["rfq_id", str(rfq_id)])
    ws.append(["rfq_ref", rfq_ref])
    ws.append(["reply_token", reply_token])
    ws.append([])
    ws.append(["pricing_row", "boq_line_item_id", "item_no"])
    for offset, item in enumerate(items):
        ws.append([FIRST_DATA_ROW + offset, str(item.boq_line_item_id), item.item_no])
    ws.sheet_state = "hidden"


def build_pricing_workbook(
    *,
    rfq_id: uuid.UUID,
    rfq_ref: str,
    reply_token: str,
    package_name: str,
    items: list[RfqLineItem],
    default_currency: str = "AED",
) -> tuple[bytes, str]:
    """Returns (xlsx_bytes, sha256_hex) -- the sha256 is stored verbatim on
    Rfq.attachment_sha256 as the durable record of what was sent.
    `default_currency` pre-fills the vendor-editable currency cell (SRS
    change #4) -- normally the tendering project's base_currency."""
    wb = Workbook()
    pricing_ws = wb.active
    assert pricing_ws is not None
    _write_pricing_sheet(pricing_ws, package_name, rfq_ref, items, default_currency)
    _write_meta_sheet(wb, rfq_id=rfq_id, rfq_ref=rfq_ref, reply_token=reply_token, items=items)

    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    return data, hashlib.sha256(data).hexdigest()
